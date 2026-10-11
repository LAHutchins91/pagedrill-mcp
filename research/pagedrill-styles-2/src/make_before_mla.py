"""Messy MLA-target sample (mla_before.docx). SAMPLE ONLY: student/instructor/course are fictional.
References are real works used as examples on MLA's own free pages; metadata checked on Crossref 2026-10-10 (crossref/)."""
import sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
out = sys.argv[1] if len(sys.argv) > 1 else "mla_before.docx"
d = Document(); s = d.sections[0]
s.left_margin = s.right_margin = Inches(1.5); s.top_margin = s.bottom_margin = Inches(0.75)
s.header.paragraphs[0].text = "Page "                           # wrong running head: label, no surname, no number
def para(t, font="Calibri", size=11, bold=False, under=False, italic_spans=(), align=None, style=None):
    p = d.add_paragraph(style=style) if style else d.add_paragraph()
    pos = 0
    for span in italic_spans:
        i = t.index(span, pos); r = p.add_run(t[pos:i]); r.font.name = font; r.font.size = Pt(size); r.bold = bold; r.underline = under
        r = p.add_run(span); r.italic = True; r.font.name = font; r.font.size = Pt(size); r.bold = bold; r.underline = under; pos = i + len(span)
    r = p.add_run(t[pos:]); r.font.name = font; r.font.size = Pt(size); r.bold = bold; r.underline = under
    if align is not None: p.alignment = align
    p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(8)
    return p
C, J = A.CENTER, A.JUSTIFY
para("Date: 14 October 2026", "Arial", 10, bold=True, align=C)
para("Student: Riley Example", "Arial", 10, bold=True, align=C)
para("Course: ENGL 102: Academic Writing", "Arial", 10, bold=True, align=C)
para("Instructor: Professor Dana Placeholder", "Arial", 10, bold=True, align=C)
para("Reading, Writing, and the Screen.", "Georgia", 16, bold=True, under=True, align=C)
para("Students now read and write on screens every day, and teachers want to know what that change means for learning.  "
     "Naomi Baron argues that reading is \u201cjust half of literacy. The other half is writing\u201d (Baron, p. 194).", "Georgia", 11, align=J)
para("1. Reading on Screens", "Cambria", 13, bold=True, align=C, style="Heading 1")
para("A long novel such as Wolf Hall can be read in print or on a screen (Mantel, 2010). "
     "Some scholars treat writing as part of what it now means to read (Baron 2013).", "Calibri", 11, italic_spans=("Wolf Hall",), align=J)
para("2. Supporting New College Writers", "Cambria", 13, bold=True, align=C, style="Heading 1")
para("Accelerated programs are one way colleges try to support basic writers (Adams et al., 2009). "
     "Second-language learning is treated as a necessity rather than an extra (Berman 12, Baron 196).", "Calibri", 11, align=J)
para("Bibliography", "Calibri", 14, bold=True, under=True)
para("Hilary Mantel - Wolf Hall (Picador 2010)", "Calibri", 11, italic_spans=("Wolf Hall",))
para("Baron, Naomi S. 2013. Redefining Reading: The Impact of Digital Communication Media. PMLA 128(1): 193-200, January.", "Calibri", 11)
para("Berman, Russell A., The Necessity of Second-Language Learning, ADFL Bulletin 43.2 (2015): 11-14. doi:10.1632/adfl.43.2.11", "Calibri", 11)
para("Peter Adams et al. The Accelerated Learning Program: Throwing Open the Gates. Journal of Basic Writing, vol 28 no 2, Fall 2009, 50-69.", "Calibri", 11)
d.core_properties.author = "Riley Example (fictional)"
d.core_properties.comments = "SAMPLE DOCUMENT: student, instructor and course are fictional. References are real works (MLA's own examples, Crossref-checked)."
d.save(out); print("wrote", out)
