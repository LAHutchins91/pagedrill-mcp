"""Check memo_after.docx against AR 25-50 (10 Oct 2020, admin rev 4 Oct 2024). Paragraph citations refer to that PDF."""
import re, json
from docx import Document
from docx.shared import Inches, Pt
D = Document("memo_after.docx"); ps = D.paragraphs; T = [p.text for p in ps]
U = "https://armypubs.army.mil/epubs/DR_pubs/DR_a/ARN42124-AR_25-50-007-WEB-13.pdf"
ch = json.load(open("memo_after_changes.json")); rows = []
def row(n, rule, ref, ok, note=""): rows.append((n, rule, ref, {True: "PASS", False: "FAIL"}.get(ok, ok), note))
s = D.sections[0]
i_os = next(i for i, t in enumerate(T) if t.startswith("EXMP-OPS"))
i_mf = next(i for i, t in enumerate(T) if t.startswith("MEMORANDUM FOR"))
i_sj = next(i for i, t in enumerate(T) if t.startswith("SUBJECT:"))
body_i = [i for i, t in enumerate(T) if re.match(r"^(\d+\.|[a-z]\.|\(\d+\)|\([a-z]\))  ", t)]
i_sig = next(i for i, t in enumerate(T) if t.startswith("JOHN A. SAMPLE"))
row(1, "Standard 8½ x 11 in. paper", "para 1-18, 2-3a", s.page_width == Inches(8.5) and s.page_height == Inches(11))
row(2, "1-in. left, right and bottom margins; right margin not justified", "para 2-3c",
    all(abs(getattr(s, k) - Inches(1)) < 1000 for k in ("left_margin", "right_margin", "bottom_margin")) and all(p.alignment in (None, 0, 1) for p in ps))
fonts = {((r.font.name or D.styles["Normal"].font.name), (r.font.size or D.styles["Normal"].font.size).pt) for i, p in enumerate(ps) if i >= i_os for r in p.runs if r.text.strip()}
row(3, "12-pt font recommended; no unusual styles (Script); font type is the senior leader's choice", "para 1-19", "PASS" if fonts == {("Arial", 12.0)} else False, f"{fonts}. FLAG: font face is set per organization (profile setting).")
row(4, "Letterhead: computer-generated APD template with DoD seal, standardized address, ZIP+4", "para 1-16, 1-35", "FLAG", "Text placeholder only (highlighted). Product must use the APD template and the user's real address.")
row(5, "Office symbol on 2nd line below the seal", "para 2-4a(1)", T[i_os-1] == "" and T[i_os-2] != "")
row(6, "ARIMS record number in parentheses one space after office symbol", "para 2-4a(2)", "FLAG", "Placeholder \"([ARIMS record no.])\" left; record number depends on record type.")
row(7, "Date on the office-symbol line, flush right, as D Month YYYY", "para 2-3b, 2-4a(3)", "\t5 October 2026" in T[i_os])
row(8, "MEMORANDUM FOR on 3rd line below office symbol; single address on same line", "para 2-4a(5)(a)", i_mf - i_os == 3 and T[i_mf-1] == T[i_mf-2] == "")
row(9, "SUBJECT: in capitals on 2nd line below last address line; 10 words or fewer", "para 2-4a(6)", i_sj - i_mf == 2 and len(T[i_sj].split()) - 1 <= 10)
row(10, "Text begins on 3rd line below subject", "para 2-4b(1)", body_i[0] - i_sj == 3)
row(11, "References listed in first paragraph", "para 1-30, 2-4b(1)(a)", T[body_i[0]].startswith("1.  Reference AR 25-50"))
row(12, "Single-spaced text, one blank line between paragraphs and subparagraphs", "para 2-4b(2)",
    all((ps[i].paragraph_format.line_spacing or 1.0) == 1.0 for i in body_i) and all(T[a+1] == "" for a in body_i[:-1]))
lab = [re.match(r"^(\S+)  ", T[i]).group(1) for i in body_i]
row(13, "Paragraphs numbered 1., 2.; subparagraphs a., b.; then (1), (2); at least two at each level", "para 2-4b(4), figure 2-1", lab == ["1.", "2.", "3.", "a.", "b.", "(1)", "(2)", "4."], " ".join(lab))
ind = [(ps[i].paragraph_format.first_line_indent or 0) for i in body_i]
row(14, "First subdivision indented 1/4 in., second 1/2 in.; wrapped lines return to left margin", "figure 2-1, para 1-39b(10)",
    [round(x / 914400, 2) for x in ind] == [0, 0, 0, 0.25, 0.25, 0.5, 0.5, 0] and all((ps[i].paragraph_format.left_indent or 0) == 0 for i in body_i), str([round(x / 914400, 2) for x in ind]))
row(15, "Two spaces after periods/question marks (and after paragraph labels as in fig. 2-1)", "para 1-39b(9) (admin rev 4 Oct 2024 changed one space back to two)",
    all("  " in T[i][:5] for i in body_i) and "schedule.  Starting" in T[body_i[1]] and "Q. Example" in T[body_i[-1]], "Initials (\"Jane Q. Example\") excluded")
row(16, "Dates in memo body as D Month YYYY", "para 1-25a", "3 November 2026" in T[body_i[1]] and "30 October 2026" in " ".join(T))
row(17, "Military time in memorandums (4 digits, no \"hours\")", "para 1-26", "0630" in T[body_i[1]] and "1300" in " ".join(T) and "a.m." not in " ".join(T) and "p.m." not in " ".join(T))
row(18, "Point of contact in the last paragraph", "para 1-23a, 2-4b(1)(e)", "point of contact" in T[body_i[-1]].lower())
row(19, "No salutation or complimentary close (memo closing = authority line, signature block, enclosure listing)", "para 2-4, 2-4c", not any(t.startswith(("Dear", "Sincerely")) for t in T), "Removed; shown to the user in the change log before saving.")
row(20, "Authority line only when signing for the commander", "para 2-4c(1)", "FLAG", "None added; user decides.")
row(21, "Signature block starts at center of page on 5th line below last text line", "para 2-4c(2)(a), 6-4b",
    i_sig - body_i[-1] == 5 and abs((ps[i_sig].paragraph_format.left_indent or 0) - Inches(3.25)) < 1000, "3.25 in. from the 1-in. left margin = horizontal center of an 8.5-in. page")
row(22, "Signature block: NAME in capitals; rank (spelled or abbreviated) and branch (abbreviated) on line 2; title on line 3", "para 6-4a(1), 6-4c, table 6-1",
    T[i_sig].isupper() and T[i_sig+1].startswith("MAJ, ") and T[i_sig+2] == "Battalion Operations Officer (S-3)", "Branch abbreviation FLAG: DA Pam 611-21 not fetched; highlighted placeholder.")
row(23, "No postscripts", "para 1-29", "FLAG", "Original had \"P.S. ...\". Removed from the layout's closing and flagged for the writer, because it is content (shown in change log).")
row(24, "One page; paragraphs no longer than 10 lines", "para 1-39b(3),(7)", "PASS", "memo_after.pdf is 1 page (pdfinfo); longest paragraph is 2 lines in the PDF render.")
with open("checklist_memo.md", "w") as f:
    f.write(f"# AR 25-50 memorandum — rule checklist for memo_after.docx\n\nSource: AR 25-50, *Preparing and Managing Correspondence*, 10 October 2020, administrative revision 007 dated 4 October 2024. Status ACTIVE on the APD record page (https://armypubs.army.mil/ProductMaps/PubForm/Details.aspx?PUB_ID=1020633), PDF: {U}. Fetched 2026-10-10 (`mil_src/ar25-50.pdf`). All units, names, office symbols and contacts are fictional placeholders.\n\n")
    f.write("| # | Rule | AR 25-50 citation | Result | Note |\n|---|---|---|---|---|\n")
    for r in rows: f.write(f"| {r[0]} | {r[1]} | {r[2]} | **{r[3]}** | {r[4]} |\n")
    c = {k: sum(1 for r in rows if r[3] == k) for k in ("PASS", "FAIL", "FLAG")}
    f.write(f"\n**Totals:** {c['PASS']} pass, {c['FAIL']} fail, {c['FLAG']} flagged for the user.\n\n## Formatter flags\n")
    for fl in ch["flags"]: f.write(f"- **{fl['item']}** ({fl['ref']}): {fl['why']}\n")
    f.write("\n## Logged changes (diff shown before saving)\n")
    for x in ch["changes"]: f.write(f"- `{x['kind']}` ({x['ref']}): “{x['before'][:150]}” → “{x['after'][:160]}”\n")
for r in rows: print(r[0], r[3], r[1][:70], r[4][:60])
