"""Prototype: lay out a messy memo as an AR 25-50 standard memorandum, driven by profiles/ar25-50_memo.json."""
import json, re, sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_COLOR_INDEX
from docx.oxml.ns import qn
prof = json.load(open("profiles/ar25-50_memo.json")); m = json.load(open("memo_extracted.json"))
orig = " ".join(p.text for p in Document("memo_before.docx").paragraphs)
def norm(s): return re.sub(r"\W+", "", s.lower())
def walk(x):
    if isinstance(x, dict): [walk(v) for k, v in x.items() if not k.startswith("_") and k != "level"]
    elif isinstance(x, list): [walk(v) for v in x]
    elif isinstance(x, str) and norm(x) not in norm(orig): raise SystemExit(f"GROUNDING FAIL: {x!r}")
walk(m)
changes, flags = [], []
def log(k, b, a, ref): changes.append({"kind": k, "before": b, "after": a, "ref": ref})
def flag(i, w, ref): flags.append({"item": i, "why": w, "ref": ref})
MONTHS = "January February March April May June July August September October November December".split()
def mil_date(s):
    mm = re.search(r"(%s)\s+(\d{1,2})(?:st|nd|rd|th)?,\s*(\d{4})" % "|".join(MONTHS), s)
    return re.sub(re.escape(mm.group(0)), f"{int(mm.group(2))} {mm.group(1)} {mm.group(3)}", s) if mm else s
def fix_text(t):
    o = t
    t = re.sub(r"(%s)\s+(\d{1,2})(?:st|nd|rd|th)?,\s*(\d{4})" % "|".join(MONTHS), lambda x: f"{int(x.group(2))} {x.group(1)} {x.group(3)}", t)
    if t != o: log("date-format", o, t, prof["date_ref"])
    o2 = t
    def mt(x):
        h, mi, ap = int(x.group(1)), int(x.group(2) or 0), x.group(3).lower()
        h = (h % 12) + (12 if ap.startswith("p") else 0)
        tail = x.string[x.end():]
        keep_period = x.group(0).endswith(".") and (tail == "" or re.match(r"\s+[A-Z]", tail))
        return f"{h:02d}{mi:02d}" + ("." if keep_period else "")
    t = re.sub(r"\b(\d{1,2})(?::(\d{2}))?\s*([ap])\.?m\.?", mt, t, flags=re.I)
    if t != o2: log("military-time", o2, t, prof["time_ref"])
    t3 = re.sub(r"(?<!\b[A-Z])([.?])\s+(?=[A-Z(])", r"\1  ", t)   # two spaces after sentence-ending period/question mark
    if t3 != t: log('sentence-spacing', t, t3, prof['spaces_ref'])
    return t3
FONT, SIZE = prof["font"]["name"], prof["font"]["size_pt"]
d = Document(); st = d.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(SIZE)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
pf = st.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
s = d.sections[0]; s.page_width, s.page_height = Inches(8.5), Inches(11)
for k, v in prof["page"]["margins_in"].items(): setattr(s, f"{k}_margin", Inches(v))
s.top_margin = Inches(0.5)   # letterhead block sits at top; see flag
def P(t="", align=WD_ALIGN_PARAGRAPH.LEFT, left=0.0, first=0.0, bold=False, size=None, hl=False):
    p = d.add_paragraph(); p.alignment = align
    p.paragraph_format.left_indent = Inches(left); p.paragraph_format.first_line_indent = Inches(first)
    if t:
        r = p.add_run(t); r.bold = bold
        if size: r.font.size = Pt(size)
        if hl: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p
def blanks(n): [P() for _ in range(n)]
C = WD_ALIGN_PARAGRAPH.CENTER
# letterhead (text placeholder; real memos use the APD computer-generated letterhead with the DoD seal)
P("[DoD SEAL — use APD letterhead template]", C, size=8, hl=True)
P("DEPARTMENT OF THE ARMY", C, bold=True, size=10)
org = m["organization"].upper(); P(org, C, bold=True, size=8)
P("[STANDARDIZED STREET ADDRESS]", C, bold=True, size=8, hl=True); P("[CITY STATE ZIP+4]", C, bold=True, size=8, hl=True)
flag("Letterhead", "AR 25-50 para 1-16b requires the computer-generated letterhead from APD's template with the DoD seal and complete standardized mailing address (ZIP+4, para 1-35). The prototype prints a highlighted text placeholder; the product must load the official template and the user's real address.", "para 1-16, 1-35")
blanks(prof["lines"]["office_symbol_after_seal"] - 1)
# office symbol + ARIMS number left, date right on same line
p = P(); p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
p.add_run(m["office_symbol"] + " ")
r = p.add_run("([ARIMS record no.])"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
date = mil_date(m["date"]); p.add_run("\t" + date)
log("date-format", m["date"], date, prof["date_ref"]); log("office-symbol-line", "FROM: EXMP-OPS (example office symbol) / Date: October 5th, 2026", f"{m['office_symbol']} ([ARIMS record no.])\t{date}", "para 2-4a(1)-(3)")
flag("ARIMS record number", "Required after the office symbol (para 2-4a(2)); depends on the record type (arims.army.mil). Not in the original, so a highlighted placeholder was left.", "para 2-4a(2)")
flag("Date", "Para 2-4a(3)(b): the date goes on after signing (or in a text box for digital signature, app F). The typed date was kept and reformatted.", "para 2-4a(3)")
blanks(prof["lines"]["memo_for_after_office_symbol"] - 1)
P("MEMORANDUM FOR " + m["memo_for"]); log("memo-for", "TO: " + m["memo_for"], "MEMORANDUM FOR " + m["memo_for"], "para 2-4a(5)(a)")
blanks(prof["lines"]["subject_after_address"] - 1)
subj = m["subject"]; subj_tc = " ".join(w if w.lower() in {"to","of","and","the","a","for","in","on"} and i else w[:1].upper() + w[1:] for i, w in enumerate(subj.split()))
P("SUBJECT: " + subj_tc)
log("subject", "Subject - " + subj, "SUBJECT: " + subj_tc, "para 2-4a(6)")
flag("Subject capitalization", "AR 25-50 requires \"SUBJECT:\" in capitals and 10 words or fewer, but the regulation text does not state a capitalization rule for the subject words. Figures show initial capitals; converted to match the figures. Confirm with local SOP.", "para 2-4a(6), figure 2-1")
blanks(prof["lines"]["text_after_subject"] - 1)
log("removed", m["salutation"], "(removed: not an element of a memorandum, para 2-4; salutations belong to letters, para 3-4)", "para 2-4, 3-4a")
# body: numbered, single-spaced, blank line between paragraphs, wrapped lines to left margin
n = {1: 0, 2: 0, 3: 0}
ind = {1: 0.0, 2: prof["body"]["first_indent_in"], 3: prof["body"]["second_indent_in"]}
paras = m["paragraphs"]
for i, para in enumerate(paras):
    L = para["level"]; n[L] += 1
    for k in range(L + 1, 4): n[k] = 0
    lab = {1: f"{n[1]}.", 2: f"{'abcdefghijklmnopqrstuvwxyz'[n[2]-1]}.", 3: f"({n[3]})"}[L]
    text = fix_text(para["text"])
    P(f"{lab}  {text}", left=0.0, first=ind[L])      # first line indented to the label; wrapped lines return to the left margin (fig 2-1)
    if i < len(paras) - 1: P()
log("numbering", "unnumbered / 'a -' / 'i.'", "1. / a. / (1) with 1/4 in. and 1/2 in. indents", prof["body"]["ref"])
flag("References paragraph", "Para 1-30 and 2-4b(1)(a): references go in the first paragraph. The original already led with the reference; kept as paragraph 1.", "para 1-30, 2-4b(1)(a)")
flag("POC", "Point of contact is in the last paragraph, as para 1-23a / 2-4b(1)(e) require. Contact details are fictional placeholders.", "para 1-23, 2-4b(1)(e)")
log("removed", m["complimentary_close"], "(removed: memorandum closing is authority line, signature block, enclosure listing, para 2-4c)", "para 2-4c")
blanks(prof["lines"]["signature_after_authority_or_text"] - 1)
flag("Authority line", "None added: only used when someone signs for the commander (para 2-4c(1)). The user must say whether one applies; signature block placed on the 5th line below the text.", "para 2-4c(1), 6-4b")
sig = m["signature"]; L0 = prof["signature"]["start_from_left_margin_in"]
name = re.sub(r"\s*\(fictional\)", "", sig["name"]).upper() + " (FICTIONAL)"
P(name, left=L0); 
p = P("", left=L0); p.add_run("MAJ, "); r = p.add_run("[branch abbrev.: Infantry]"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
P(sig["title"], left=L0)
log("signature-block", f"{sig['name']} / {sig['rank']}, {sig['branch']} / {sig['title']}", f"{name} / MAJ, [branch abbrev.] / {sig['title']}", "para 6-4a-c, table 6-1")
flag("Branch abbreviation", "Para 6-4c says branch must be abbreviated, per DA Pam 611-21 (not fetched in this prototype). Left highlighted instead of guessing the code.", "para 6-4c, 6-4f")
flag("Postscript", f"\"{m['postscript']}\" — para 1-29: do not use postscripts in Army correspondence. It is the writer's content, so it is held out of the layout and listed in the change log for approval; the writer can cut it or move the thought into the body.", "para 1-29")
log("held-out", m["postscript"], "(held out of the memo pending the writer's OK: no postscripts, para 1-29)", "para 1-29")
flag("Font", "Para 1-19 leaves font type to Army senior leaders (12 pt recommended). Arial 12 used as a placeholder profile setting; set the unit's font in the profile.", "para 1-19")
flag("Paragraph length / one page", "Para 1-39b(3),(7): paragraphs of no more than 10 lines; one-page memos. Both met here (checked visually in the PDF).", "para 1-39b")
d.core_properties.comments = "SAMPLE: all units, names, office symbols and contact details are fictional placeholders."
d.save("memo_after.docx")
json.dump({"changes": changes, "flags": flags}, open("memo_after_changes.json", "w"), indent=2, ensure_ascii=False)
print(len(changes), "changes", len(flags), "flags")
