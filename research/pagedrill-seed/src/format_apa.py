"""Prototype: reformat a messy student paper to APA 7 student format, driven by a style profile.
Rules: format only. Never rewrites the author's sentences, never invents reference data.
Every text-level change (labels, citation punctuation, heading case/numbering) is logged in changes.json
so it can be shown to the user as a diff before saving (Stetwise-style propose -> apply)."""
import json, re, sys, copy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, DST, PROFILE, REFS = sys.argv[1:5]
prof = json.load(open(PROFILE)); refs = json.load(open(REFS))["references"]
FONT, SIZE = prof["font"]["name"], prof["font"]["size_pt"]
changes, flags = [], []
def log(kind, before, after, rule): changes.append({"kind": kind, "before": before, "after": after, "rule": rule})
def flag(item, why, rule): flags.append({"item": item, "why": why, "rule": rule})

src = Document(SRC)
paras = [(p.style.name, p.text.strip()) for p in src.paragraphs if p.text.strip()]

# ---------- 1. classify roles ----------
meta, body, ref_texts, title = {}, [], [], None
LABELS = {"name": "author", "author": "author", "class": "course", "course": "course",
          "professor": "instructor", "instructor": "instructor", "teacher": "instructor",
          "due": "due_date", "date": "due_date", "institution": "affiliation", "school": "affiliation"}
in_refs = False
for style, text in paras:
    m = re.match(r"^(\w+)\s*:\s*(.+)$", text)
    if not in_refs and m and m.group(1).lower() in LABELS and not body:
        meta[LABELS[m.group(1).lower()]] = m.group(2).strip(); continue
    if re.match(r"^(works cited|references?|bibliography|reference list)\s*:?$", text, re.I):
        in_refs = True; log("label", text, prof["references"]["label"], prof["references"]["src"]); continue
    if in_refs: ref_texts.append(text); continue
    if title is None and not style.startswith("Heading") and not body:
        title = text; continue
    lvl = int(style.split()[-1]) if style.startswith("Heading") and style.split()[-1].isdigit() else 0
    body.append({"level": lvl, "text": text})

# ---------- 2. title case (APA definition) ----------
MINOR = {"a","an","the","and","as","but","for","if","nor","or","so","yet","at","by","in","of","off","on","per","to","up","via"}
def title_case(s):
    words = s.split(" "); out = []; cap_next = True
    for w in words:
        core = w.lower()
        bare = re.sub(r"[^\w-]", "", core)
        if cap_next or bare not in MINOR or len(bare) > 3:
            w2 = "-".join(x[:1].upper() + x[1:] for x in core.split("-"))
        else: w2 = core
        out.append(w2); cap_next = bool(re.search(r"[:.?!\u2014]$", w))
    return " ".join(out)

def fix_heading(t, lvl):
    orig = t
    t2 = re.sub(r"^\s*(\d+(\.\d+)*[.)]|[a-zA-Z][.)]|[IVXivx]+[.)])\s+", "", t)
    if t2 != t: log("heading-numbering", t, t2, prof["headings"]["src"])
    was_caps = t2.isupper()
    t3 = title_case(t2)
    if t3 != t2: log("heading-case", t2, t3, prof["headings"]["case_src"])
    if was_caps: flag(f'Heading "{orig}"', "Was ALL CAPS; converted to title case. Check acronyms/proper nouns.", prof["headings"]["case_src"])
    return t3

if title:
    nt = title_case(title)
    if nt != title: log("title-case", title, nt, prof["headings"]["case_src"]); title = nt

# drop "Introduction" heading at start (APA: paper title acts as de facto Level 1 heading)
if body and body[0]["level"] and body[0]["text"].strip().lower().strip(":") in ("introduction", "intro"):
    log("removed-heading", body[0]["text"], "(removed; paper title acts as Level 1 heading)", prof["headings"]["src"]); body.pop(0)
for b in body:
    if b["level"]: b["text"] = fix_heading(b["text"], b["level"])

# one-subsection check
def check_single_subsections():
    for i, b in enumerate(body):
        if not b["level"]: continue
        L = b["level"]; kids = []
        for c in body[i+1:]:
            if c["level"] and c["level"] <= L: break
            if c["level"] == L + 1: kids.append(c["text"])
        if len(kids) == 1:
            flag(f'Section "{b["text"]}"', f'Has only one Level {L+1} subsection ("{kids[0]}"). APA: avoid a single subsection. Needs the writer to add a second subsection or drop the heading; not changed automatically.', prof["headings"]["src"])
check_single_subsections()

# ---------- 3. references: grounding check + render ----------
def norm(s): return re.sub(r"[\W_]+", "", s.lower())
def ground(ref, original):
    o = norm(original)
    for k, v in ref.items():
        if k in ("key", "type") or v is None: continue
        vals = [a.get("family") for a in v] + [a.get("given") or "" for a in v] if k == "authors" else [v]
        for x in vals:
            if x and norm(x) not in o: raise SystemExit(f"GROUNDING FAIL: {k}={x!r} not in original reference text")
orig_for = {}
for r in refs:
    hit = [t for t in ref_texts if norm(r["authors"][0]["family"]) in norm(t)]
    orig_for[r["key"]] = hit[0]; ground(r, hit[0])

def initials(a): return a.get("initials") or " ".join(p[0] + "." for p in re.split(r"[\s-]+", a["given"]) if p)
def ref_authors(a):
    names = [f'{x["family"]}, {initials(x)}' for x in a]
    if len(names) == 1: return names[0]
    return ", ".join(names[:-1]) + ", & " + names[-1]
def dash(p): return p.replace("-", "\u2013") if p else p

rendered = []   # list of (sortkey, [(text, italic, highlight)])
for r in refs:
    segs = [(ref_authors(r["authors"]) + " ", False, False)]
    if r["type"] == "journal-article":
        segs.append((f'({r["year"]}). ', False, False))
        segs.append((r["title"].rstrip(".") + ("" if r["title"].endswith("?") else ".") + " ", False, False))
        segs.append((r["journal"], True, False)); segs.append((", ", False, False))
        segs.append((r["volume"], True, False))
        if r.get("issue"): segs.append((f'({r["issue"]})', False, False))
        else: flag(f'Reference {r["authors"][0]["family"]} ({r["year"]})', "No issue number in the original. APA says always include the issue number for a journal article when the journal uses one; if this journal/volume has no issues, omitting is correct. Not added (no guessing).", "https://apastyle.apa.org/style-grammar-guidelines/references/examples/journal-article-references")
        segs.append((f', {dash(r["pages"])}.', False, False))
        if r.get("doi"):
            d = re.sub(r"^(https?://(dx\.)?doi\.org/|doi:\s*)", "", r["doi"], flags=re.I)
            segs.append((f" https://doi.org/{d}", False, False))
            if "doi:" in orig_for[r["key"]].lower(): log("doi-format", "doi:" + d, "https://doi.org/" + d, prof["references"]["doi_src"])
        else: flag(f'Reference {r["authors"][0]["family"]} ({r["year"]})', "No DOI in the original. If the article has a DOI, APA requires it; the formatter does not look one up or add it. (Papers by Ouroboros could look it up for the user to confirm.)", prof["references"]["doi_src"])
    elif r["type"] == "webpage-news":
        if r.get("date"): segs.append((f'({r["date"]}). ', False, False))
        else:
            segs.append(("(", False, False)); segs.append(("[DATE NEEDED]", False, True)); segs.append(("). ", False, False))
            flag(f'Reference {r["authors"][0]["family"]}', "No publication date in the original. APA wants the specific date for news webpages; use \"n.d.\" only if the page truly has none. Left a highlighted [DATE NEEDED] marker instead of guessing. (The URL contains 2020/06/04, but URL dates are not reliable: APA's own example for this same story gives 2020, September 24.)", prof["references"]["missing_src"])
        segs.append((r["title"].rstrip(".") + ".", True, False)); segs.append((f' {r["site"]}. {r["url"]}', False, False))
    sk = (r["authors"][0]["family"].lower())
    rendered.append((sk, segs, r))
    log("reference", orig_for[r["key"]], "".join(s[0] for s in segs), prof["references"]["elements_src"])
before_order = [r["key"] for r in refs]
rendered.sort(key=lambda x: x[0])
if [x[2]["key"] for x in rendered] != before_order:
    log("reference-order", ", ".join(before_order), ", ".join(x[2]["key"] for x in rendered), prof["references"]["src"])

# ---------- 4. in-text citations ----------
REF = {r["authors"][0]["family"]: r for r in refs}
def year_of(r): return r.get("year") or (r.get("date") or "")[:4] or None
def paren_form(r, yr):
    a = r["authors"]; fam = [x["family"] for x in a]
    nm = fam[0] if len(fam) == 1 else (f"{fam[0]} & {fam[1]}" if len(fam) == 2 else f"{fam[0]} et al.")
    return f"{nm}, {yr}"
def narr_form(r, yr):
    fam = [x["family"] for x in r["authors"]]
    nm = fam[0] if len(fam) == 1 else (f"{fam[0]} and {fam[1]}" if len(fam) == 2 else f"{fam[0]} et al.")
    return f"{nm} ({yr})"

def fix_citations(text):
    orig = text
    # parenthetical groups
    def paren(m):
        inner = m.group(1); items = [i.strip() for i in inner.split(";")]
        out = []; changed = False
        for it in items:
            fam = next((f for f in REF if re.match(rf"^{re.escape(f)}\b", it)), None)
            if not fam: return m.group(0)
            ym = re.search(r"\b(1[89]\d\d|20\d\d|n\.d\.)[a-z]?\b", it)
            if not ym:
                flag(f'Citation "({it})"', "No year in the citation, and the reference has no date either. APA needs author + year. Left as typed.", prof["citations"]["src"])
                out.append(it); continue
            out.append(paren_form(REF[fam], ym.group(1)))
        if len(out) > 1: out = sorted(out, key=lambda s: s.lower())
        return "(" + "; ".join(out) + ")"
    text = re.sub(r"\(([^()]*[A-Z][a-z]+[^()]*)\)", paren, text)
    # narrative: "A, B, and C (YEAR)" | "A and B YEAR" | "A and B (YEAR)"
    for fam, r in REF.items():
        names = [x["family"] for x in r["authors"]]
        if len(names) >= 3: pat = re.escape(names[0]) + r",\s*" + r",?\s*".join(map(re.escape, names[1:-1])) + r",?\s*(?:and|&)\s*" + re.escape(names[-1])
        elif len(names) == 2: pat = re.escape(names[0]) + r"\s*(?:and|&)\s*" + re.escape(names[1])
        else: pat = re.escape(names[0])
        text = re.sub(pat + r"\s*\(?\s*((?:19|20)\d\d)\)?(?![,;])", lambda m: narr_form(r, m.group(1)), text)
    if text != orig: log("in-text-citation", orig, text, prof["citations"]["src"])
    return text
for b in body:
    if not b["level"]: b["text"] = fix_citations(b["text"])

# ---------- 5. build the document ----------
out = Document()
st = out.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(SIZE)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
pf = st.paragraph_format; pf.line_spacing = prof["line_spacing"]["value"]; pf.space_before = Pt(0); pf.space_after = Pt(0)
for lvl in range(1, 6):     # restyle built-in headings so Word's navigation pane still works (accessibility advice on headings page)
    hs = out.styles[f"Heading {lvl}"]; spec = prof["headings"][str(lvl)]
    hs.font.name = FONT; hs.font.size = Pt(SIZE); hs.font.bold = spec["bold"]; hs.font.italic = spec["italic"]
    hs.font.color.rgb = RGBColor(0, 0, 0); hs.font.underline = False
    rpr = hs.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    hp = hs.paragraph_format; hp.space_before = Pt(0); hp.space_after = Pt(0); hp.line_spacing = 2.0
    hp.first_line_indent = Inches(0); hp.left_indent = Inches(0); hp.keep_with_next = True
sec = out.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("top", "bottom", "left", "right"): setattr(sec, f"{side}_margin", Inches(prof["page"]["margins_in"][side]))
sec.header_distance = Inches(0.5)

# header: page number top right, every page, title page = 1; no running head for student papers
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
def field(p, instr):
    r = p.add_run(); r.font.name = FONT; r.font.size = Pt(SIZE)
    for tag, attr in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None), ("w:fldChar", {"w:fldCharType": "separate"})):
        e = OxmlElement(tag)
        if attr:
            for k, v in attr.items(): e.set(qn(k), v)
        else: e.set(qn("xml:space"), "preserve"); e.text = f" {instr} "
        r._r.append(e)
    t = OxmlElement("w:t"); t.text = "1"; r._r.append(t)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)
field(hp, "PAGE")
log("page-header", src.sections[0].header.paragraphs[0].text, "[page number, right-aligned]", prof["page_header"]["src"])

def P(text="", align="left", bold=False, italic=False, indent=True, style=None, hanging=None):
    p = out.add_paragraph(style=style)
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    f = p.paragraph_format; f.line_spacing = 2.0; f.space_before = Pt(0); f.space_after = Pt(0)
    f.first_line_indent = Inches(prof["paragraph"]["first_line_indent_in"]) if indent else Inches(0)
    if hanging: f.left_indent = Inches(hanging); f.first_line_indent = Inches(-hanging)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.name = FONT; r.font.size = Pt(SIZE)
    return p

# title page
tp = prof["title_page"]
for _ in range(tp["blank_lines_before_title"]): P(indent=False)
P(title, "center", bold=True, indent=False)
P(indent=False)                        # one double-spaced blank line between title and byline
course = meta.get("course")
if course and not re.match(r"^[A-Z]{2,5}\s?\d{2,4}[A-Z]?:", course):
    c2 = re.sub(r"^([A-Z]{2,5}\s?\d{2,4}[A-Z]?)\s+", r"\1: ", course)
    if c2 != course: log("course-format", course, c2, tp["src"]); course = c2
flag("Course name", f'Kept as typed ("{course}"). APA: give the course number and name as shown on instructional materials; confirm the official course name.', tp["src"])
for key in ("author", "affiliation", "course", "instructor", "due_date"):
    val = course if key == "course" else meta.get(key)
    if val: P(val, "center", indent=False)
    else:
        p = P("", "center", indent=False); r = p.add_run(f"[{key.replace('_', ' ').upper()} NEEDED: Department, Institution]" if key == "affiliation" else f"[{key.upper()} NEEDED]")
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        flag(f"Title page: {key}", "Not in the original paper. APA's student title page requires it (e.g., \"Department of Psychology, University of Georgia\"). Left a highlighted placeholder; the formatter never invents a school.", tp["src"])
for lab, k in (("Name:", "author"), ("Professor:", "instructor"), ("Due:", "due_date")):
    if meta.get(k): log("title-page-label-removed", f"{lab} {meta[k]}", meta[k], tp["src"])
out.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

# first page of text: title, bold, centered
P(title, "center", bold=True, indent=False)
i = 0
while i < len(body):
    b = body[i]; L = b["level"]
    if L:
        spec = prof["headings"][str(L)]
        if spec["inline"]:
            nxt = body[i+1]["text"] if i + 1 < len(body) and not body[i+1]["level"] else ""
            p = P("", "left", indent=True, style=f"Heading {L}")
            r = p.add_run(b["text"].rstrip(".") + "."); r.bold = True; r.italic = spec["italic"]
            r2 = p.add_run(" " + nxt); r2.bold = False; r2.italic = False
            i += 2 if nxt else 1; continue
        P(b["text"], "center" if spec["align"] == "center" else "left", bold=spec["bold"], italic=spec["italic"], indent=False, style=f"Heading {L}")
    else:
        P(b["text"], "left", indent=True)
    i += 1

# references on a new page
out.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
lab = prof["references"]["label"] if len(rendered) != 1 else "Reference"
P(lab, "center", bold=True, indent=False)
for _, segs, r in rendered:
    p = P("", "left", indent=False, hanging=prof["references"]["hanging_indent_in"])
    for text, ital, hl in segs:
        run = p.add_run(text); run.italic = ital
        if hl: run.font.highlight_color = WD_COLOR_INDEX.YELLOW
out.core_properties.author = "Jordan Sample (fictional)"; out.core_properties.title = title
out.core_properties.comments = "SAMPLE DOCUMENT: student, instructor and course are fictional placeholders. References are real works."
out.save(DST)
json.dump({"changes": changes, "flags": flags}, open(DST.replace(".docx", "_changes.json"), "w"), indent=2, ensure_ascii=False)
print(f"wrote {DST}: {len(changes)} logged text/format changes, {len(flags)} flags")
