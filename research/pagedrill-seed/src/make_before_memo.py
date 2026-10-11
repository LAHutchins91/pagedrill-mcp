"""Messy Army memo (memo_before.docx). EVERYTHING here is a fictional placeholder: unit, office symbol,
names, phone (555-01xx fictional range) and email (example.com reserved domain)."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
d = Document(); s = d.sections[0]
for k in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(s, k, Inches(0.6))
def P(t, size=11, font="Times New Roman", bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=None):
    p = d.add_paragraph(); r = p.add_run(t); r.font.size = Pt(size); r.font.name = font; r.bold = bold
    p.alignment = align; p.paragraph_format.line_spacing = 2.0
    if indent is not None: p.paragraph_format.first_line_indent = Inches(indent)
    return p
P("1st Example Battalion (FICTIONAL UNIT - EXAMPLE ONLY)", 14, "Old English Text MT", True, WD_ALIGN_PARAGRAPH.LEFT)
P("Date: October 5th, 2026", align=WD_ALIGN_PARAGRAPH.LEFT)
P("TO: Commander, Example Company A (EXMP-A)", align=WD_ALIGN_PARAGRAPH.LEFT)
P("FROM: EXMP-OPS (example office symbol)", align=WD_ALIGN_PARAGRAPH.LEFT)
P("Subject - change to physical readiness training schedule", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
P("Dear Commander,", align=WD_ALIGN_PARAGRAPH.LEFT)
P("Reference AR 25-50 (Preparing and Managing Correspondence).", indent=0.5)
P("This memorandum announces a change to the battalion physical readiness training schedule. Starting November 3, 2026, "
  "morning training will begin at 6:30 a.m. instead of 6:00 a.m.", indent=0.5)
P("Company commanders will do the following:", indent=0.5)
P("a - Update the company training calendar.", indent=1.0)
P("b - Brief all Soldiers at the next formation.", indent=1.0)
P("i. Include the new start time in the brief.", indent=1.5)
P("ii. Collect questions and send them to the S-3 by 1:00 p.m. on October 30, 2026.", indent=1.5)
P("The point of contact for this memorandum is CPT Jane Q. Example (fictional), (555) 555-0100, jane.q.example@example.com.", indent=0.5)
P("Sincerely,", align=WD_ALIGN_PARAGRAPH.LEFT)
P("John A. Sample (fictional)", align=WD_ALIGN_PARAGRAPH.LEFT)
P("Major, Infantry", align=WD_ALIGN_PARAGRAPH.LEFT)
P("Battalion Operations Officer (S-3)", align=WD_ALIGN_PARAGRAPH.LEFT)
P("P.S. Thanks for your help with this change.", align=WD_ALIGN_PARAGRAPH.LEFT)
d.core_properties.comments = "SAMPLE: all units, names, office symbols and contact details are fictional placeholders."
d.save("memo_before.docx"); print("wrote memo_before.docx")
