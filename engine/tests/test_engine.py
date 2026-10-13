"""Sample documents against the verified checklists."""

import base64
import unittest
from pathlib import Path

from docx import Document
from io import BytesIO

from pagedrill.bibliography import render_imported
from pagedrill.engine import check_loaded, dispatch, format_loaded
from pagedrill.errors import PagedrillRefusal
from pagedrill.profiles import get_profile, require_selectable
from pagedrill.words import assert_apa_body_words

ROOT = Path(__file__).resolve().parents[2]


def load(name: str) -> str:
    return base64.b64encode((ROOT / "samples" / name).read_bytes()).decode("ascii")


def direct_format(style_id: str, data_b64: str, supplied: dict | None = None) -> dict:
    """Exercise a verified engine while public dispatch still refuses a hidden style."""
    try:
        result = format_loaded(get_profile(style_id), base64.b64decode(data_b64), supplied or {})
        return {"ok": True, "result": result}
    except PagedrillRefusal as exc:
        return {"ok": False, "error": {"code": exc.code, "reason": str(exc)}}


def direct_check(style_id: str, data_b64: str) -> dict:
    return {"ok": True, "result": check_loaded(get_profile(style_id), base64.b64decode(data_b64))}


def direct_render(style_id: str, kind: str, text: str) -> dict:
    return {"ok": True, "result": render_imported(text, kind, get_profile(style_id))}


class Styles(unittest.TestCase):
    def test_only_verified_styles_are_selectable(self):
        styles = dispatch({"action": "list_styles"})["result"]["styles"]
        selectable = {item["id"] for item in styles if item["selectable"]}
        self.assertEqual(
            selectable,
            {"apa7-student", "apa7-professional", "army-ar25-50-memorandum"},
        )
        listed = {item["id"] for item in styles}
        self.assertEqual(listed, selectable)
        for style_id in ("mla9-research-paper", "ieee-conference-letter", "vancouver-icmje-nlm"):
            self.assertNotIn(style_id, listed)
            with self.assertRaises(PagedrillRefusal) as caught:
                require_selectable(style_id)
            self.assertEqual(caught.exception.code, "style_not_selectable")
            refused = dispatch({"action": "format", "style_id": style_id, "data_b64": load("apa_before.docx")})
            self.assertFalse(refused["ok"])
            self.assertEqual(refused["error"]["code"], "style_not_selectable")

    def test_apa_student_checklist_and_word_equality(self):
        result = dispatch({"action": "format", "style_id": "apa7-student", "data_b64": load("apa_before.docx")})
        self.assertTrue(result["ok"], result.get("error"))
        counts = result["result"]["checklist"]["counts"]
        self.assertEqual(counts["FAIL"], 0)
        self.assertEqual(counts["PASS"], 28)
        flags = [row["id"] for row in result["result"]["checklist"]["rules"] if row["result"] == "FLAG"]
        self.assertEqual(flags, [11, 12, 18, 27, 31])
        self.assertTrue(result["result"]["verification"]["untouched_identical"])
        self.assertEqual(result["result"]["verification"]["added_parts"], [])
        document = Document(BytesIO(base64.b64decode(result["result"]["data_b64"])))
        blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertNotIn("doi:", blob.lower())
        self.assertIn("https://doi.org/10.1017/S0144686X20001932", blob)
        self.assertIn("[DATE NEEDED]", blob)
        self.assertNotIn("10.1016/j.chb.2017.02.038", blob)
        self.assertIn("[AFFILIATION NEEDED", blob)

    def test_user_supplied_values_are_used_and_not_invented(self):
        result = dispatch(
            {
                "action": "format",
                "style_id": "apa7-student",
                "data_b64": load("apa_before.docx"),
                "supplied": {
                    "affiliation": "Department of Psychology, Example University",
                    "references": [{"family": "Sanchiz", "doi": "10.1016/j.chb.2017.02.038"}],
                },
            }
        )
        self.assertTrue(result["ok"], result.get("error"))
        document = Document(BytesIO(base64.b64decode(result["result"]["data_b64"])))
        blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Department of Psychology, Example University", blob)
        self.assertIn("https://doi.org/10.1016/j.chb.2017.02.038", blob)
        self.assertIn("[DATE NEEDED]", blob)

    def test_professional_running_head_is_not_invented(self):
        result = dispatch({"action": "format", "style_id": "apa7-professional", "data_b64": load("apa_before.docx")})
        self.assertTrue(result["ok"], result.get("error"))
        self.assertEqual(result["result"]["checklist"]["counts"]["FAIL"], 0)
        document = Document(BytesIO(base64.b64decode(result["result"]["data_b64"])))
        header = document.sections[0].header.paragraphs[0].text
        self.assertIn("[RUNNING HEAD NEEDED]", header)
        self.assertNotIn("Running head:", header)
        blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertNotIn("PSY 201", blob)
        self.assertIn("[AUTHOR NOTE NEEDED]", blob)
        self.assertIn("Author Note", blob)

    def test_memo_checklist(self):
        result = dispatch({"action": "format", "style_id": "army-ar25-50-memorandum", "data_b64": load("memo_before.docx")})
        self.assertTrue(result["ok"], result.get("error"))
        counts = result["result"]["checklist"]["counts"]
        self.assertEqual(counts, {"PASS": 20, "FAIL": 0, "FLAG": 4})
        self.assertIn("Unclassified", result["result"]["notice"])
        document = Document(BytesIO(base64.b64decode(result["result"]["data_b64"])))
        blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertNotIn("Thanks for your help", blob)
        self.assertNotIn("Dear Commander", blob)
        self.assertIn("[ARIMS record no.]", blob)
        self.assertIn("[branch abbrev.: Infantry]", blob)
        self.assertTrue(result["result"]["verification"]["untouched_identical"])
        self.assertNotIn("word/header1.xml", result["result"]["verification"]["changed_parts"])

    def test_messy_check_does_not_modify_and_reports_failures(self):
        original = (ROOT / "samples" / "apa_before.docx").read_bytes()
        result = dispatch({"action": "check", "style_id": "apa7-student", "data_b64": base64.b64encode(original).decode("ascii")})
        self.assertTrue(result["ok"])
        self.assertGreater(result["result"]["checklist"]["counts"]["FAIL"], 0)
        self.assertEqual(len(result["result"]["checklist"]["rules"]), 33)
        self.assertNotIn("data_b64", result["result"])

    def test_word_change_is_refused(self):
        with self.assertRaises(PagedrillRefusal) as caught:
            assert_apa_body_words("The cat sat.", "The dog sat.", set(), set())
        self.assertEqual(caught.exception.code, "word_change")

    def test_unverified_style_is_refused(self):
        result = dispatch({"action": "format", "style_id": "mla9-research-paper", "data_b64": load("apa_before.docx")})
        self.assertFalse(result["ok"])
        self.assertEqual(result["error"]["code"], "style_not_selectable")

    def test_ieee_conference_checklist_and_journal_references(self):
        result = direct_format("ieee-conference-letter", load("ieee_before.docx"))
        self.assertTrue(result["ok"], result.get("error"))
        counts = result["result"]["checklist"]["counts"]
        self.assertEqual(counts, {"PASS": 28, "FAIL": 0, "FLAG": 7})
        self.assertEqual(result["result"]["verification"]["basis"], "ieee-conference-template")
        document = Document(BytesIO(base64.b64decode(result["result"]["data_b64"])))
        blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn(
            "M. M. Chiampi and L. L. Zilberti, \u201cInduction of electric field in human bodies moving near MRI: An efficient BEM computational procedure,\u201d IEEE Trans. Biomed. Eng., vol. 58, no. 10, pp. 2787\u20132793, Oct. 2011, doi: 10.1109/TBME.2011.2158315.",
            blob,
        )
        original = (ROOT / "samples" / "ieee_before.docx").read_bytes()
        checked = direct_check("ieee-conference-letter", base64.b64encode(original).decode("ascii"))
        self.assertTrue(checked["ok"], checked.get("error"))
        self.assertGreater(checked["result"]["checklist"]["counts"]["FAIL"], 0)
        self.assertEqual(len(checked["result"]["checklist"]["rules"]), 35)
        self.assertNotIn("data_b64", checked["result"])

    def test_nonjournal_reference_is_kept_and_flagged(self):
        source = Document(BytesIO((ROOT / "samples" / "ieee_before.docx").read_bytes()))
        book = "Smith, J. Display Hardware. New York: Example Press, 2019."
        source.add_paragraph(book)
        raw = BytesIO()
        source.save(raw)
        result = direct_format("ieee-conference-letter", base64.b64encode(raw.getvalue()).decode("ascii"))
        self.assertTrue(result["ok"], result.get("error"))
        self.assertEqual(result["result"]["checklist"]["counts"]["FAIL"], 0)
        document = Document(BytesIO(base64.b64decode(result["result"]["data_b64"])))
        blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn(book, blob)
        reasons = " ".join(item.get("why", "") for item in result["result"]["flags"])
        self.assertIn("not a journal article", reasons)

    def test_tracked_changes_and_toc_keep_the_checklist(self):
        import zipfile

        result = dispatch(
            {
                "action": "format",
                "style_id": "apa7-student",
                "data_b64": load("apa_before.docx"),
                "supplied": {"track_changes": True, "table_of_contents": True},
            }
        )
        self.assertTrue(result["ok"], result.get("error"))
        self.assertEqual(result["result"]["checklist"]["counts"], {"PASS": 28, "FAIL": 0, "FLAG": 5})
        package = zipfile.ZipFile(BytesIO(base64.b64decode(result["result"]["data_b64"])))
        xml = package.read("word/document.xml")
        self.assertIn(b'w:author="Pagedrill"', xml)
        self.assertIn(b"<w:ins", xml)
        self.assertIn(b"<w:del ", xml)
        self.assertIn(b"TOC", xml)

    def test_imported_references_do_not_invent_fields(self):
        bib = """@article{smith2020,
  author = {Smith, Ada B.},
  title = {A short title about displays},
  journal = {Applied Physics Letters},
  year = {2020},
  volume = {10}
}
@book{jones2020,
  author = {Jones, Pat},
  title = {A Book of Displays},
  year = {2020},
  publisher = {Example Press}
}
"""
        result = direct_render("ieee-conference-letter", "bib", bib)
        self.assertTrue(result["ok"], result.get("error"))
        entries = result["result"]["entries"]
        article, book = entries
        rendered = article["rendered"]
        self.assertIsInstance(rendered, str)
        self.assertNotIn("doi", rendered.lower())
        self.assertNotIn("pp.", rendered)
        self.assertNotIn("May", rendered)
        self.assertTrue(any(name in article["missing"] for name in ("doi", "pages", "issue", "month")))
        self.assertIsNone(book["rendered"])
        self.assertIn("A Book of Displays", str(book["fields"]))
        ris = "TY  - JOUR\nAU  - Smith, Ada B.\nTI  - A short title about displays\nJO  - Applied Physics Letters\nPY  - 2020\nVL  - 10\nER  - \n"
        apa = dispatch({"action": "render_references", "style_id": "apa7-student", "format": "ris", "text": ris})
        self.assertTrue(apa["ok"], apa.get("error"))
        sentence = apa["result"]["entries"][0]["rendered"]
        self.assertIn("Smith, A. B. (2020). A short title about displays. Applied Physics Letters, 10.", sentence)
        self.assertNotIn("doi", sentence.lower())
        self.assertNotIn("https://", sentence)

    def test_create_memorandum_flags_letterhead_and_arims(self):
        result = dispatch(
            {
                "action": "create_memo",
                "fields": {
                    "organization": "Example Battalion",
                    "office_symbol": "EXMP-OPS",
                    "memo_for": "Commander, Example Battalion",
                    "subject": "Range Schedule for the Drill",
                    "paragraphs": ["Report at 0630 on 30 October 2026. Bring water."],
                    "signature": {"name": "Alex Quinn", "rank": "Major", "title": "Operations Officer", "branch": "Infantry"},
                    "date": "October 5, 2026",
                    "poc": "The point of contact is Alex Quinn.",
                },
            }
        )
        self.assertTrue(result["ok"], result.get("error"))
        self.assertEqual(result["result"]["checklist"]["counts"]["FAIL"], 0)
        self.assertEqual(result["result"]["verification"]["basis"], "new-memorandum")
        document = Document(BytesIO(base64.b64decode(result["result"]["data_b64"])))
        blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("[DoD SEAL", blob)
        self.assertIn("[ARIMS record no.]", blob)
        self.assertIn("EXAMPLE BATTALION", blob)
        self.assertIn("MEMORANDUM FOR Commander, Example Battalion", blob)
        self.assertIn("SUBJECT: Range Schedule for the Drill", blob)
        self.assertIn("Report at 0630 on 30 October 2026", blob)
        self.assertIn("Bring water.", blob)
        self.assertIn("The point of contact is Alex Quinn.", blob)
        self.assertIn("ALEX QUINN", blob)
        self.assertIn("MAJ", blob)
        self.assertIn("[branch abbrev.: Infantry]", blob)
        self.assertIn("5 October 2026", blob)
        flagged = " ".join(f"{item.get('item', '')} {item.get('why', '')}" for item in result["result"]["flags"]).lower()
        self.assertIn("letterhead", flagged)
        self.assertIn("arims", flagged)
        again = dispatch(
            {"action": "format", "style_id": "army-ar25-50-memorandum", "data_b64": result["result"]["data_b64"]}
        )
        self.assertTrue(again["ok"], again.get("error"))
        self.assertEqual(again["result"]["checklist"]["counts"]["FAIL"], 0)


class Vancouver(unittest.TestCase):
    def test_sample_matches_nlm_and_citing_medicine(self):
        from pagedrill.vancouver import OFFICIAL, PAGE_EXAMPLES, elide_pages, render_parsed
        from pagedrill.profiles import get_profile

        for source, expected in PAGE_EXAMPLES.items():
            self.assertEqual(elide_pages(source), expected, source)
        result = direct_format("vancouver-icmje-nlm", load("vanc_before.docx"))
        self.assertTrue(result["ok"], result.get("error"))
        self.assertEqual(result["result"]["checklist"]["counts"], {"PASS": 12, "FAIL": 0, "FLAG": 5})
        self.assertEqual(result["result"]["verification"]["basis"], "in-place")
        self.assertEqual(result["result"]["verification"]["changed_parts"], ["word/document.xml"])
        document = Document(BytesIO(base64.b64decode(result["result"]["data_b64"])))
        blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for entry in OFFICIAL.values():
            self.assertIn(entry, blob)
        self.assertIn("(1, 4)", blob)
        self.assertNotIn("Halpern et al.", blob)
        flagged = " ".join(item.get("why", "") for item in result["result"]["flags"]).lower()
        self.assertIn("journal option", flagged)
        self.assertIn("proper noun", flagged)
        self.assertIn("12088834", flagged)
        original = (ROOT / "samples" / "vanc_before.docx").read_bytes()
        checked = direct_check("vancouver-icmje-nlm", base64.b64encode(original).decode("ascii"))
        self.assertTrue(checked["ok"], checked.get("error"))
        self.assertGreater(checked["result"]["checklist"]["counts"]["FAIL"], 0)
        self.assertEqual(len(checked["result"]["checklist"]["rules"]), 17)

        profile = get_profile("vancouver-icmje-nlm")
        authors = [{"family": f"Author{index}", "initials": "A. B."} for index in range(7)]
        book, _flags, _missing = render_parsed(
            {
                "type": "book",
                "authors": authors,
                "key": "author02000",
                "year": "2000",
                "title": "A Long Textbook",
                "edition": "2nd",
                "place": "Paris",
                "publisher": "Example Press",
                "typed": "A Long Textbook",
            },
            profile,
        )
        self.assertIsInstance(book, str)
        self.assertNotIn("et al", book)
        self.assertEqual(book.count("Author"), 7)

    def test_journal_options_are_not_guessed(self):
        import zipfile

        raw = load("vanc_before.docx")
        comma = direct_format("vancouver-icmje-nlm", raw, {"journal_options": {"multiple_citation": "comma"}})
        self.assertTrue(comma["ok"], comma.get("error"))
        self.assertEqual(comma["result"]["checklist"]["counts"]["FAIL"], 0)
        document = Document(BytesIO(base64.b64decode(comma["result"]["data_b64"])))
        blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("(1,4)", blob)
        self.assertNotIn("(1, 4)", blob)
        tracked = direct_format("vancouver-icmje-nlm", raw, {"track_changes": True})
        self.assertTrue(tracked["ok"], tracked.get("error"))
        self.assertEqual(tracked["result"]["checklist"]["counts"], {"PASS": 12, "FAIL": 0, "FLAG": 5})
        package = zipfile.ZipFile(BytesIO(base64.b64decode(tracked["result"]["data_b64"])))
        xml = package.read("word/document.xml")
        self.assertIn(b'w:author="Pagedrill"', xml)
        self.assertIn(b"<w:ins", xml)
        refused = direct_format("vancouver-icmje-nlm", raw, {"journal_options": {"page_layout": "times"}})
        self.assertFalse(refused["ok"])
        self.assertEqual(refused["error"]["code"], "bad_request")
        bib = """@article{ada2020,
  author = {Smith, Ada B.},
  title = {A short title about displays},
  journal = {Applied Physics Letters},
  year = {2020},
  volume = {10},
  pages = {284-287}
}
@book{murray2002,
  author = {Murray, P. R. and Rosenthal, K. S. and Kobayashi, G. S. and Pfaller, M. A.},
  title = {Medical Microbiology},
  year = {2002},
  edition = {4th},
  address = {St. Louis},
  publisher = {Mosby}
}
"""
        imported = direct_render("vancouver-icmje-nlm", "bib", bib)
        self.assertTrue(imported["ok"], imported.get("error"))
        article, book = imported["result"]["entries"]
        self.assertIn("Applied Physics Letters", article["rendered"])
        self.assertIn("284-7", article["rendered"])
        self.assertNotIn("doi", article["rendered"].lower())
        self.assertIn("doi", article["missing"])
        self.assertEqual(
            book["rendered"],
            "Murray PR, Rosenthal KS, Kobayashi GS, Pfaller MA. Medical microbiology. 4th ed. St. Louis: Mosby; 2002.",
        )


if __name__ == "__main__":
    unittest.main()
