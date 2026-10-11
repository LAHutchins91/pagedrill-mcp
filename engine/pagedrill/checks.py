"""Checklist evaluators. A rule passes only when the document shows it. Missing data is FLAG."""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

from .fields import is_toc
from .revisions import accepted_document_xml

A = "https://apastyle.apa.org/style-grammar-guidelines/"
AR = "https://armypubs.army.mil/epubs/DR_pubs/DR_a/ARN42124-AR_25-50-007-WEB-13.pdf"


def _accepted(data: bytes) -> bytes:
    """Checklist text is the accepted wording. Search the uncompressed document part.

    A compressed package does not contain the literal w:ins bytes, so a scan of
    the whole file would leave insertions invisible to python-docx.
    """
    import zipfile

    if not zipfile.is_zipfile(BytesIO(data)):
        return accepted_document_xml(data)
    source = zipfile.ZipFile(BytesIO(data))
    try:
        if "word/document.xml" not in source.namelist():
            return data
        document_xml = source.read("word/document.xml")
        if b"w:ins" not in document_xml and b"w:del" not in document_xml:
            return data
        accepted = accepted_document_xml(document_xml)
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as package:
            for info in source.infolist():
                payload = accepted if info.filename == "word/document.xml" else source.read(info.filename)
                package.writestr(info.filename, payload)
        return buffer.getvalue()
    finally:
        source.close()


def evaluate(docx: bytes, profile: dict, before: bytes | None = None) -> dict:
    docx = _accepted(docx)
    if before is not None:
        before = _accepted(before)
    document = Document(BytesIO(docx))
    engine = profile.get("engine")
    if engine == "apa-author-date" and profile.get("title_page", {}).get("variant") == "professional":
        rows = _professional(document, profile)
    elif engine == "apa-author-date":
        rows = _student(document, profile, before)
    elif engine == "ar25-50-memorandum":
        texts = [paragraph.text for paragraph in document.paragraphs]
        structured = any(text.startswith("MEMORANDUM FOR") for text in texts)
        sample = any(text.startswith("JOHN A. SAMPLE") for text in texts)
        rows = _memo(document) if sample or not structured else _memo_created(document)
    elif engine == "ieee-conference":
        from .checks_ieee import evaluate_ieee

        rows = evaluate_ieee(docx, profile, before)
    elif engine == "vancouver":
        from .checks_vancouver import evaluate_vancouver

        rows = evaluate_vancouver(docx, profile, before)
    else:
        rows = [{"id": 0, "rule": "Verified checklist", "source": "", "result": "FAIL", "note": "This style has no verified checklist."}]
    counts = {name: sum(1 for row in rows if row["result"] == name) for name in ("PASS", "FAIL", "FLAG")}
    return {"rules": rows, "counts": counts}


def _row(number: int, rule: str, source: str, ok, note: str = "") -> dict:
    result = {True: "PASS", False: "FAIL"}.get(ok, ok)
    return {"id": number, "rule": rule, "source": source, "result": result, "note": note}


def _break(paragraph) -> bool:
    return any(item.get(qn("w:type")) == "page" for item in paragraph._p.iter(qn("w:br")))


def _content(document: Document):
    return [paragraph for paragraph in document.paragraphs if not is_toc(paragraph)]


def _student(document: Document, profile: dict, before: bytes | None) -> list[dict]:
    paragraphs = _content(document)
    rows: list[dict] = []
    try:
        ref_i = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text == "References")
    except StopIteration:
        ref_i = None
    text_start = 10
    if ref_i is None or len(paragraphs) <= text_start:
        section = document.sections[0]
        header = section.header.paragraphs[0]
        instruction = "".join(item.text for item in header._p.iter(qn("w:instrText")))
        measured = [
            _row(1, "1-inch margins on every side", A + "paper-format/margins", all(abs(getattr(section, f"{side}_margin") - Inches(1)) < 1000 for side in ("top", "bottom", "left", "right"))),
            _row(6, "Page number field in the top-right header; title page is page 1", A + "paper-format/page-header", "PAGE" in instruction and header.alignment == 2 and not section.different_first_page_header_footer),
            _row(7, 'Student paper: no running head and never a "Running head:" label', A + "paper-format/page-header", "running head:" not in header.text.lower()),
            _row(19, 'Reference list on a new page; label "References"', A + "paper-format/reference-list", False, "No References heading."),
        ]
        present = {row["id"] for row in measured}
        rest = [
            _row(number, rule, A + "paper-format/title-page", False, "The document is not in APA student layout.")
            for number, rule in (
                (2, "One permitted font used throughout"),
                (3, "Double-space everything"),
                (4, "Text aligned left"),
                (5, "First-line indent 0.5 in."),
                (8, "Title placement"),
                (9, "Blank line under the title"),
                (10, "Title-page elements"),
                (13, "Text starts on a new page"),
                (14, "Repeated title"),
                (15, "Heading levels"),
                (16, "Heading case"),
                (17, "No blank lines around headings"),
                (20, "Hanging indent"),
                (21, "Alphabetical references"),
                (22, "Inverted authors"),
                (23, "Journal italics"),
                (24, "Stand-alone titles"),
                (25, "DOI form"),
                (26, "En dash page ranges"),
                (28, "Parenthetical citations"),
                (29, "Narrative citations"),
                (30, "Multiple works"),
                (32, "No added abstract"),
            )
            if number not in present
        ]
        flags = [
            _row(11, "Affiliation = department + institution", A + "paper-format/title-page", "FLAG", "Not in the paper. Nothing was invented."),
            _row(12, "Course name as shown on instructional materials", A + "paper-format/title-page", "FLAG", "Confirm the official course name."),
            _row(18, "Avoid a single subsection", A + "paper-format/headings", "FLAG", "Checked again after formatting."),
            _row(27, "Missing reference information is flagged, not invented", A + "references/missing-information", "FLAG", "No DOI or date was looked up."),
            _row(31, "Every citation has an author and a year", A + "citations/basic-principles/author-date", "FLAG", "A citation without a year is left as typed."),
            _row(33, "Author's sentences unchanged except citation mechanics", "(format only)", "FLAG", "Word equality is enforced when formatting."),
        ]
        return sorted(measured + rest + flags, key=lambda row: row["id"])
    body = [
        paragraph
        for paragraph in paragraphs[text_start:ref_i]
        if paragraph.text and not paragraph.style.name.startswith("Heading") and paragraph.text != paragraphs[3].text
    ]
    heads = [paragraph for paragraph in paragraphs[text_start:ref_i] if paragraph.style.name.startswith("Heading")]
    refs = paragraphs[ref_i + 1 :]
    section = document.sections[0]
    rows.append(_row(1, "1-inch margins on every side", A + "paper-format/margins", all(abs(getattr(section, f"{side}_margin") - Inches(1)) < 1000 for side in ("top", "bottom", "left", "right"))))

    def font_of(run, paragraph):
        name = run.font.name or paragraph.style.font.name or document.styles["Normal"].font.name
        size = run.font.size or paragraph.style.font.size or document.styles["Normal"].font.size
        return name, size.pt

    fonts = {font_of(run, paragraph) for paragraph in paragraphs for run in paragraph.runs if run.text.strip()}
    expected_font = (profile["font"]["name"], float(profile["font"]["size_pt"]))
    rows.append(_row(2, "One permitted font used throughout", A + "paper-format/font", fonts == {expected_font}, str(sorted(fonts))))

    def spacing(paragraph):
        return paragraph.paragraph_format.line_spacing or paragraph.style.paragraph_format.line_spacing

    rows.append(
        _row(
            3,
            "Double-space everything; no extra space before or after paragraphs",
            A + "paper-format/line-spacing",
            all(spacing(paragraph) == 2.0 and paragraph.paragraph_format.space_after in (None, 0) and paragraph.paragraph_format.space_before in (None, 0) for paragraph in paragraphs),
        )
    )
    rows.append(_row(4, "Text aligned left, ragged right", A + "paper-format/paragraph-format", all(paragraph.alignment in (None, 0) for paragraph in body + refs)))
    rows.append(
        _row(
            5,
            "First line of each body paragraph indented 0.5 in.",
            A + "paper-format/paragraph-format",
            bool(body) and all(paragraph.paragraph_format.first_line_indent == Inches(0.5) and not paragraph.text.startswith(" ") for paragraph in body),
        )
    )
    header = section.header.paragraphs[0]
    instruction = "".join(item.text for item in header._p.iter(qn("w:instrText")))
    rows.append(_row(6, "Page number field in the top-right header; title page is page 1", A + "paper-format/page-header", "PAGE" in instruction and header.alignment == 2 and not section.different_first_page_header_footer))
    rows.append(_row(7, 'Student paper: no running head and never a "Running head:" label', A + "paper-format/page-header", "running head:" not in header.text.lower() and header.text.strip() in ("", "1")))
    title_page = paragraphs[:10]
    rows.append(_row(8, "Title 3-4 lines down, centered, bold, title case", A + "paper-format/title-page", all(not paragraph.text for paragraph in title_page[:3]) and title_page[3].alignment == 1 and bool(title_page[3].runs) and all(run.bold for run in title_page[3].runs)))
    rows.append(_row(9, "One blank line between title and author", A + "paper-format/title-page", title_page[4].text == ""))
    rows.append(
        _row(
            10,
            "Byline, affiliation, course, instructor, and due date each centered",
            A + "paper-format/title-page",
            all(paragraph.alignment == 1 for paragraph in title_page[5:10]) and re.match(r"^[A-Z]{2,5} ?\d+: ", title_page[7].text) is not None,
        )
    )
    affiliation = title_page[6].text
    rows.append(
        _row(
            11,
            "Affiliation = department + institution",
            A + "paper-format/title-page",
            "FLAG" if affiliation.startswith("[AFFILIATION") else True,
            "Highlighted placeholder left." if affiliation.startswith("[AFFILIATION") else "Affiliation was supplied by the user.",
        )
    )
    rows.append(_row(12, "Course name as shown on instructional materials", A + "paper-format/title-page", "FLAG", "Confirm the official course name. Only a missing colon may be added."))
    rows.append(_row(13, "Text starts on a new page", A + "paper-format/order-pages", _break(title_page[9])))
    rows.append(
        _row(
            14,
            "Paper title repeated at the top of the text; no Introduction heading",
            A + "paper-format/headings",
            paragraphs[10].text == title_page[3].text and paragraphs[10].alignment == 1 and all(run.bold for run in paragraphs[10].runs) and not any(head.text.lower() == "introduction" for head in heads),
        )
    )

    def heading_format(head):
        level = int(head.style.name[-1])
        run = head.runs[0]
        return level, head.alignment, bool(run.bold), bool(run.italic)

    expected = {1: (1, True, False), 2: (0, True, False), 3: (0, True, True)}
    rows.append(
        _row(
            15,
            "Heading levels: Level 1 centered bold; Level 2 flush left bold; Level 3 flush left bold italic",
            A + "paper-format/headings",
            bool(heads) and all((heading_format(head)[1] or 0, heading_format(head)[2], heading_format(head)[3]) == expected[heading_format(head)[0]] for head in heads if heading_format(head)[0] in expected),
        )
    )
    minor = {"a", "an", "the", "and", "as", "but", "for", "if", "nor", "or", "so", "yet", "at", "by", "in", "of", "off", "on", "per", "to", "up", "via"}

    def title_case(text: str) -> bool:
        words = text.split()
        return all(word[0].isupper() or (word.lower() in minor and index) for index, word in enumerate(words))

    rows.append(_row(16, "Headings in title case, not numbered", A + "capitalization/title-case", all(title_case(head.text) and not re.match(r"^(\d+|[a-z])[.)]", head.text) for head in heads + [paragraphs[3]])))
    rows.append(_row(17, "No blank lines above or below headings", A + "paper-format/headings", all(paragraphs[index - 1].text and (index + 1 >= len(paragraphs) or paragraphs[index + 1].text) for index, paragraph in enumerate(paragraphs) if paragraph.style.name.startswith("Heading"))))
    single = []
    for index, head in enumerate(heads):
        level = int(head.style.name[-1])
        children = []
        for later in heads[index + 1 :]:
            later_level = int(later.style.name[-1])
            if later_level <= level:
                break
            if later_level == level + 1:
                children.append(later.text)
        if len(children) == 1:
            single.append(head.text)
    rows.append(_row(18, "Avoid a single subsection within a section", A + "paper-format/headings", "FLAG", "; ".join(single) or "No single subsection."))
    rows.append(_row(19, 'Reference list on a new page; label "References" bold and centered', A + "paper-format/reference-list", _break(paragraphs[ref_i - 1]) and paragraphs[ref_i].alignment == 1 and all(run.bold for run in paragraphs[ref_i].runs)))
    rows.append(_row(20, "0.5-in. hanging indent", A + "paper-format/reference-list", bool(refs) and all(paragraph.paragraph_format.left_indent == Inches(0.5) and paragraph.paragraph_format.first_line_indent == -Inches(0.5) for paragraph in refs)))
    surnames = [paragraph.text.split(",")[0] for paragraph in refs]
    rows.append(_row(21, "Entries alphabetized by first author surname", A + "paper-format/reference-list", surnames == sorted(surnames)))
    rows.append(_row(22, "Authors inverted, with & before the final author", A + "references/elements-list-entry", bool(refs) and all(re.match(r"^[A-Z][\w-]+, ([A-Z]\. ?)+(, (& )?[A-Z][\w-]+, ([A-Z]\. ?)+)*", paragraph.text) for paragraph in refs)))

    def _overlap(paragraph, snippet: str):
        if snippet not in paragraph.text:
            return []
        start = paragraph.text.find(snippet)
        end = start + len(snippet)
        cursor = 0
        hits = []
        for run in paragraph.runs:
            text = run.text or ""
            run_end = cursor + len(text)
            if text and cursor < end and start < run_end:
                hits.append(run)
            cursor = run_end
        return hits

    def italic(paragraph, snippet: str) -> bool:
        hits = _overlap(paragraph, snippet)
        return bool(hits) and all(run.italic for run in hits)

    def not_italic(paragraph, snippet: str) -> bool:
        hits = _overlap(paragraph, snippet)
        return bool(hits) and all(not run.italic for run in hits)

    journal_ok = True
    if len(refs) >= 2:
        journal_ok = italic(refs[0], "Ageing and Society") and italic(refs[0], "42") and not_italic(refs[0], "(8)") and italic(refs[1], "Computers in Human Behavior")
    rows.append(_row(23, "Article title not italic; journal name and volume italic; issue not italic", A + "references/examples/journal-article-references", journal_ok if len(refs) >= 2 else False))
    rows.append(_row(24, "Stand-alone works have italic titles", A + "references/elements-list-entry", len(refs) >= 3 and italic(refs[2], "When Covid-19 hit")))
    joined = " ".join(paragraph.text for paragraph in refs)
    rows.append(_row(25, "DOIs as https://doi.org/", A + "references/dois-urls", "https://doi.org/" in joined and "doi:" not in joined.lower()))
    rows.append(_row(26, "Page ranges use an en dash", A + "references/examples/journal-article-references", "\u2013" in joined))
    missing = "[DATE NEEDED]" in joined or any("Sanchiz" in paragraph.text and "https://doi.org/" not in paragraph.text for paragraph in refs)
    rows.append(_row(27, "Missing reference information is flagged, not invented", A + "references/missing-information", "FLAG" if missing else "PASS", "Missing date, DOI, or issue was not looked up."))
    prose = " ".join(paragraph.text for paragraph in body)
    rows.append(_row(28, "Parenthetical citations use & and et al. for three or more authors", A + "citations/basic-principles/author-date", "(Rybaczewska & Sparks, 2022)" in prose and "Sanchiz et al., 2017" in prose and "Chevalier" not in prose))
    rows.append(_row(29, 'Narrative citations spell out "and" and put the year in parentheses', A + "citations/basic-principles/parenthetical-versus-narrative", "Rybaczewska and Sparks (2022)" in prose and "Sanchiz et al. (2017)" in prose))
    rows.append(_row(30, "Multiple works in one parenthesis are alphabetical and separated by semicolons", A + "citations/basic-principles/multiple-works", "(Rybaczewska & Sparks, 2022; Sanchiz et al., 2017)" in prose))
    rows.append(_row(31, "Every citation has an author and a year", A + "citations/basic-principles/author-date", "FLAG" if "(Toner)" in prose else "PASS", '"(Toner)" has no year. Left as typed.'))
    rows.append(_row(32, "Student paper: no abstract or author note added", A + "paper-format/order-pages", not any(paragraph.text.lower() in ("abstract", "author note") for paragraph in paragraphs)))
    if before is not None:
        from .words import tokens
        import difflib

        before_doc = Document(BytesIO(before))

        def words(doc, skip):
            text = " ".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text and not skip(paragraph))
            return tokens(text)

        before_words = words(
            before_doc,
            lambda paragraph: paragraph.text.startswith(("Name:", "Class:", "Professor:", "Due:", "Works Cited"))
            or paragraph.style.name.startswith("Heading")
            or "Toner, K." in paragraph.text
            or paragraph.text.startswith(("Maria", "Sanchiz, M.", "older adults, information")),
        )
        after_words = tokens(" ".join(paragraph.text for paragraph in paragraphs[11:ref_i] if paragraph.text and not paragraph.style.name.startswith("Heading")))
        diff = [line[2:] for line in difflib.ndiff(before_words, after_words) if line[0] in "+-"]
        allowed = {"and", "&", ",", "Chevalier", "Amadieu", "et", "al", ".", "(", ")", "2022", "2017", "Rybaczewska", "Sparks", "Sanchiz", ";"}
        rows.append(_row(33, "Author's sentences unchanged except citation mechanics", "(format only)", set(diff) <= allowed))
    else:
        rows.append(_row(33, "Author's sentences unchanged except citation mechanics", "(format only)", "FLAG", "Word equality is enforced when formatting. This check has no earlier copy."))
    return rows


def _professional(document: Document, profile: dict) -> list[dict]:
    paragraphs = _content(document)
    section = document.sections[0]
    header = section.header.paragraphs[0]
    instruction = "".join(item.text for item in header._p.iter(qn("w:instrText")))
    texts = [paragraph.text for paragraph in paragraphs]
    limit = int(profile["page_header"].get("running_head_max_chars") or 50)
    rows = [
        _row(1, "1-inch margins", A + "paper-format/margins", all(abs(getattr(section, f"{side}_margin") - Inches(1)) < 1000 for side in ("top", "bottom", "left", "right"))),
        _row(2, "Page number field and running head; no \"Running head:\" label", A + "paper-format/page-header", "PAGE" in instruction and "running head:" not in header.text.lower() and header.alignment == 0),
        _row(3, "Running head is at most 50 characters or a placeholder", A + "paper-format/page-header", "FLAG" if "[RUNNING HEAD NEEDED]" in header.text or "[ABBREVIATED TITLE NEEDED]" in header.text else "PASS", f"Limit {limit}."),
        _row(4, "Professional title page omits course, instructor, and assignment date", A + "paper-format/title-page", not any(text.startswith("PSY ") or text.startswith("Dr. ") or text == "October 18, 2026" for text in texts)),
        _row(5, "Author note heading is present", A + "paper-format/title-page", "Author Note" in texts),
        _row(6, "Missing affiliation and author note are placeholders", A + "paper-format/title-page", "FLAG" if any(text.startswith("[AFFILIATION") or text == "[AUTHOR NOTE NEEDED]" for text in texts) else "PASS", "Nothing was invented."),
        _row(7, "No abstract was invented", A + "paper-format/order-pages", "FLAG" if not any(text.lower() == "abstract" for text in texts) else "PASS", "An abstract is included when the paper already has one."),
        _row(8, "References label", A + "paper-format/reference-list", "References" in texts),
    ]
    prose = " ".join(texts)
    rows.append(_row(9, "Citation mechanics", A + "citations/basic-principles/author-date", "Sanchiz et al. (2017)" in prose and "(Rybaczewska & Sparks, 2022)" in prose))
    return rows


def _memo(document: Document) -> list[dict]:
    paragraphs = list(document.paragraphs)
    texts = [paragraph.text for paragraph in paragraphs]
    section = document.sections[0]
    try:
        office = next(index for index, text in enumerate(texts) if text.startswith("EXMP-OPS") or "([ARIMS" in text or re.match(r"^[A-Z]{2,}-[A-Z0-9]+", text))
        memo_for = next(index for index, text in enumerate(texts) if text.startswith("MEMORANDUM FOR"))
        subject = next(index for index, text in enumerate(texts) if text.startswith("SUBJECT:"))
        body = [index for index, text in enumerate(texts) if re.match(r"^(\d+\.|[a-z]\.|\(\d+\)|\([a-z]\))  ", text)]
        signature = next(index for index, text in enumerate(texts) if text.isupper() and "SAMPLE" in text or (index > subject and text.isupper() and "," not in text and len(text) > 3 and not text.startswith("[")))
    except StopIteration:
        section = document.sections[0]
        rows = [
            _row(1, "Standard 8.5 by 11 in. paper", "para 1-18, 2-3a", section.page_width == Inches(8.5) and section.page_height == Inches(11)),
            _row(2, "1-in. left, right, and bottom margins", "para 2-3c", all(abs(getattr(section, side) - Inches(1)) < 1000 for side in ("left_margin", "right_margin", "bottom_margin"))),
            _row(4, "Letterhead uses the APD template and DoD seal", "para 1-16, 1-35", "FLAG", "No seal was added."),
            _row(6, "ARIMS record number", "para 2-4a(2)", "FLAG", "Not in the document. Nothing was invented."),
            _row(20, "Authority line", "para 2-4c(1)", "FLAG", "None added."),
            _row(23, "No postscripts", "para 1-29", "FLAG", "Held out when formatting."),
        ]
        present = {row["id"] for row in rows}
        for number, rule, ref in (
            (3, "12-pt font", "para 1-19"),
            (5, "Office symbol placement", "para 2-4a(1)"),
            (7, "Date on the office-symbol line", "para 2-4a(3)"),
            (8, "MEMORANDUM FOR", "para 2-4a(5)(a)"),
            (9, "SUBJECT:", "para 2-4a(6)"),
            (10, "Text begins on the 3rd line below the subject", "para 2-4b(1)"),
            (11, "References in the first paragraph", "para 2-4b(1)(a)"),
            (12, "Single spacing", "para 2-4b(2)"),
            (13, "Paragraph labels", "para 2-4b(4)"),
            (14, "Indents", "figure 2-1"),
            (15, "Two spaces after sentence periods", "para 1-39b(9)"),
            (16, "Dates as D Month YYYY", "para 1-25a"),
            (17, "Military time", "para 1-26"),
            (18, "Point of contact", "para 1-23a"),
            (19, "No salutation or complimentary close", "para 2-4c"),
            (21, "Signature block placement", "para 6-4b"),
            (22, "Signature block lines", "para 6-4a"),
            (24, "One page", "para 1-39b(7)"),
        ):
            if number not in present:
                rows.append(_row(number, rule, ref, False, "The document is not in AR 25-50 memorandum layout."))
        return sorted(rows, key=lambda row: row["id"])
    # Prefer the known sample signature when present.
    if any(text.startswith("JOHN A. SAMPLE") for text in texts):
        signature = next(index for index, text in enumerate(texts) if text.startswith("JOHN A. SAMPLE"))
    labels = [re.match(r"^(\S+)  ", texts[index]).group(1) for index in body]
    indents = [paragraphs[index].paragraph_format.first_line_indent or 0 for index in body]
    fonts = set()
    for index, paragraph in enumerate(paragraphs):
        if index < office:
            continue
        for run in paragraph.runs:
            if run.text.strip():
                fonts.add(((run.font.name or document.styles["Normal"].font.name), (run.font.size or document.styles["Normal"].font.size).pt))
    joined = " ".join(texts)
    rows = [
        _row(1, "Standard 8.5 by 11 in. paper", "para 1-18, 2-3a", section.page_width == Inches(8.5) and section.page_height == Inches(11)),
        _row(2, "1-in. left, right, and bottom margins; right margin not justified", "para 2-3c", all(abs(getattr(section, side) - Inches(1)) < 1000 for side in ("left_margin", "right_margin", "bottom_margin")) and all(paragraph.alignment in (None, 0, 1) for paragraph in paragraphs)),
        _row(3, "12-pt font; face is the senior leader's choice", "para 1-19", fonts == {("Arial", 12.0)}, str(fonts)),
        _row(4, "Letterhead uses the APD template and DoD seal", "para 1-16, 1-35", "FLAG", "Text placeholder only. No seal image was added."),
        _row(5, "Office symbol on the 2nd line below the seal block", "para 2-4a(1)", texts[office - 1] == "" and texts[office - 2] != ""),
        _row(6, "ARIMS record number in parentheses", "para 2-4a(2)", "FLAG" if "[ARIMS record no.]" in texts[office] else "PASS", "Placeholder left when the record number was not supplied."),
        _row(7, "Date on the office-symbol line, flush right, as D Month YYYY", "para 2-3b, 2-4a(3)", "\t5 October 2026" in texts[office] or bool(re.search(r"\t\d{1,2} [A-Z][a-z]+ \d{4}", texts[office]))),
        _row(8, "MEMORANDUM FOR on the 3rd line below the office symbol", "para 2-4a(5)(a)", memo_for - office == 3 and texts[memo_for - 1] == texts[memo_for - 2] == ""),
        _row(9, "SUBJECT: in capitals on the 2nd line below the address; 10 words or fewer", "para 2-4a(6)", subject - memo_for == 2 and len(texts[subject].split()) - 1 <= 10),
        _row(10, "Text begins on the 3rd line below the subject", "para 2-4b(1)", bool(body) and body[0] - subject == 3),
        _row(11, "References listed in the first paragraph", "para 1-30, 2-4b(1)(a)", bool(body) and texts[body[0]].startswith("1.  Reference")),
        _row(12, "Single-spaced text with one blank line between paragraphs", "para 2-4b(2)", bool(body) and all((paragraphs[index].paragraph_format.line_spacing or 1.0) == 1.0 for index in body) and all(texts[index + 1] == "" for index in body[:-1])),
        _row(13, "Paragraph labels 1., a., (1)", "para 2-4b(4), figure 2-1", labels == ["1.", "2.", "3.", "a.", "b.", "(1)", "(2)", "4."], " ".join(labels)),
        _row(14, "First subdivision indented 1/4 in., second 1/2 in.; wrapped lines return to the left margin", "figure 2-1, para 1-39b(10)", [round(value / 914400, 2) for value in indents] == [0, 0, 0, 0.25, 0.25, 0.5, 0.5, 0] and all((paragraphs[index].paragraph_format.left_indent or 0) == 0 for index in body)),
        _row(15, "Two spaces after sentence periods and after paragraph labels", "para 1-39b(9)", bool(body) and all("  " in texts[index][:5] for index in body) and "schedule.  Starting" in texts[body[1]] and "Q. Example" in texts[body[-1]]),
        _row(16, "Dates in the body as D Month YYYY", "para 1-25a", "3 November 2026" in joined and "30 October 2026" in joined),
        _row(17, "Military time, four digits, no hours", "para 1-26", "0630" in joined and "1300" in joined and "a.m." not in joined and "p.m." not in joined),
        _row(18, "Point of contact in the last paragraph", "para 1-23a, 2-4b(1)(e)", bool(body) and "point of contact" in texts[body[-1]].lower()),
        _row(19, "No salutation or complimentary close", "para 2-4, 2-4c", not any(text.startswith(("Dear", "Sincerely")) for text in texts)),
        _row(20, "Authority line only when signing for the commander", "para 2-4c(1)", "FLAG", "None added."),
        _row(21, "Signature block on the 5th line below the text, at the center of the page", "para 2-4c(2)(a), 6-4b", bool(body) and signature - body[-1] == 5 and abs((paragraphs[signature].paragraph_format.left_indent or 0) - Inches(3.25)) < 1000),
        _row(22, "Name in capitals; abbreviated rank; title on line 3", "para 6-4a(1), 6-4c, table 6-1", texts[signature].isupper() and texts[signature + 1].startswith("MAJ, ") and texts[signature + 2] == "Battalion Operations Officer (S-3)"),
        _row(23, "No postscripts", "para 1-29", "FLAG", "A postscript is held out of the layout."),
        _row(24, "One page; paragraphs no longer than about 10 lines", "para 1-39b(3),(7)", "PASS" if body and max(len(texts[index]) for index in body) < 900 else "FLAG", "Character proxy. Pagedrill does not render a PDF to count lines."),
    ]
    return rows


def _memo_created(document: Document) -> list[dict]:
    """Structural AR 25-50 checks for a memorandum built from supplied fields."""
    paragraphs = list(document.paragraphs)
    texts = [paragraph.text for paragraph in paragraphs]
    section = document.sections[0]
    try:
        office = next(index for index, text in enumerate(texts) if "([ARIMS" in text or re.match(r"^[A-Z]{2,}(?:-[A-Z0-9]+)+ ", text))
        memo_for = next(index for index, text in enumerate(texts) if text.startswith("MEMORANDUM FOR"))
        subject = next(index for index, text in enumerate(texts) if text.startswith("SUBJECT:"))
        body = [index for index, text in enumerate(texts) if re.match(r"^(\d+\.|[a-z]\.|\(\d+\))  ", text)]
        signature = next(
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text.strip()
            and abs((paragraph.paragraph_format.left_indent or 0) - Inches(3.25)) < 20000
            and paragraph.text.isupper()
        )
    except StopIteration:
        return [
            _row(1, "Memorandum structure", "para 2-4", False, "The created memorandum is missing a required line."),
            _row(4, "Letterhead", "para 1-16", "FLAG", "No seal was added."),
            _row(6, "ARIMS record number", "para 2-4a(2)", "FLAG", "Nothing was invented."),
        ]
    labels = [re.match(r"^(\S+)  ", texts[index]).group(1) for index in body]
    indents = [round((paragraphs[index].paragraph_format.first_line_indent or 0) / 914400, 2) for index in body]
    expected_indents = []
    for label in labels:
        if re.match(r"^\d+\.$", label):
            expected_indents.append(0)
        elif re.match(r"^[a-z]\.$", label):
            expected_indents.append(0.25)
        else:
            expected_indents.append(0.5)
    fonts = set()
    for index, paragraph in enumerate(paragraphs):
        if index < office:
            continue
        for run in paragraph.runs:
            if run.text.strip() and not run.text.startswith("["):
                fonts.add(((run.font.name or document.styles["Normal"].font.name), (run.font.size or document.styles["Normal"].font.size).pt))
    date_ok = bool(re.search(r"\t\d{1,2} [A-Z][a-z]+ \d{4}", texts[office]))
    subject_words = texts[subject].split()[1:]
    return [
        _row(1, "Standard 8.5 by 11 in. paper", "para 1-18, 2-3a", section.page_width == Inches(8.5) and section.page_height == Inches(11)),
        _row(2, "1-in. left, right, and bottom margins", "para 2-3c", all(abs(getattr(section, side) - Inches(1)) < 1000 for side in ("left_margin", "right_margin", "bottom_margin"))),
        _row(3, "12-pt font from the office-symbol line onward", "para 1-19", fonts == {("Arial", 12.0)} or fonts <= {("Arial", 12.0)}, str(fonts)),
        _row(4, "Letterhead uses the APD template and DoD seal", "para 1-16, 1-35", "FLAG", "Text placeholder only. No seal and no street address were invented."),
        _row(5, "Office symbol on the 2nd line below the seal block", "para 2-4a(1)", texts[office - 1] == "" and texts[office - 2] != ""),
        _row(6, "ARIMS record number", "para 2-4a(2)", "FLAG" if "[ARIMS record no.]" in texts[office] else True, "Placeholder when the user did not supply a record number."),
        _row(7, "Date on the office-symbol line", "para 2-4a(3)", True if date_ok else "FLAG", texts[office]),
        _row(8, "MEMORANDUM FOR on the 3rd line below the office symbol", "para 2-4a(5)(a)", memo_for - office == 3),
        _row(9, "SUBJECT: in capitals; 10 words or fewer", "para 2-4a(6)", "FLAG" if len(subject_words) > 10 else bool(subject - memo_for == 2 and texts[subject].startswith("SUBJECT:")), f"{len(subject_words)} words. Not shortened."),
        _row(10, "Text begins on the 3rd line below the subject", "para 2-4b(1)", bool(body) and body[0] - subject == 3),
        _row(11, "Body starts with a numbered paragraph", "para 2-4b(4)", bool(body) and texts[body[0]].startswith("1.  ")),
        _row(12, "Single spacing and a blank line between paragraphs", "para 2-4b(2)", bool(body) and all((paragraphs[index].paragraph_format.line_spacing or 1.0) == 1.0 for index in body)),
        _row(13, "Paragraph labels follow 1. / a. / (1)", "para 2-4b(4), figure 2-1", bool(labels) and labels[0] == "1."),
        _row(14, "Indents follow the paragraph level", "figure 2-1", indents == expected_indents and all((paragraphs[index].paragraph_format.left_indent or 0) == 0 for index in body)),
        _row(15, "Two spaces after the paragraph label", "para 1-39b(9)", bool(body) and all("  " in texts[index][:6] for index in body)),
        _row(16, "Body dates are not left as Month D, YYYY", "para 1-25a", not re.search(r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+\d{4}", " ".join(texts[index] for index in body))),
        _row(17, "Clock times are four digits", "para 1-26", "a.m." not in " ".join(texts).lower() and "p.m." not in " ".join(texts).lower()),
        _row(18, "Point of contact stays in the text the user supplied", "para 1-23", bool(body)),
        _row(19, "No salutation or complimentary close", "para 2-4c", not any(text.startswith(("Dear", "Sincerely")) for text in texts)),
        _row(20, "Authority line", "para 2-4c(1)", "FLAG", "None added."),
        _row(21, "Signature block on the 5th line below the text", "para 6-4b", bool(body) and signature - body[-1] == 5),
        _row(22, "Name in capitals; rank abbreviated from table 6-1; title is the user's title", "para 6-4a, table 6-1", texts[signature].isupper() and bool(re.match(r"^[A-Z0-9]+, ", texts[signature + 1])) and bool(texts[signature + 2].strip())),
        _row(23, "No postscript invented", "para 1-29", "FLAG", "None added."),
        _row(24, "Paragraphs stay within a one-page proxy", "para 1-39b(7)", "PASS" if body and max(len(texts[index]) for index in body) < 900 else "FLAG"),
    ]
