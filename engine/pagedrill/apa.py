"""APA 7 student and professional papers. Layout comes from the profile. Missing data is flagged."""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .errors import PagedrillRefusal
from .fields import add_toc_field
from .packageio import splice_allowed
from .revisions import apply_tracked_changes
from .words import assert_apa_body_words

MINOR = {"a", "an", "the", "and", "as", "but", "for", "if", "nor", "or", "so", "yet", "at", "by", "in", "of", "off", "on", "per", "to", "up", "via"}
LABELS = {
    "name": "author",
    "author": "author",
    "class": "course",
    "course": "course",
    "professor": "instructor",
    "instructor": "instructor",
    "teacher": "instructor",
    "due": "due_date",
    "date": "due_date",
    "institution": "affiliation",
    "school": "affiliation",
}


def _norm(value: str) -> str:
    return re.sub(r"[\W_]+", "", value.lower())


def title_case(value: str) -> str:
    words = value.split(" ")
    out: list[str] = []
    cap_next = True
    for word in words:
        core = word.lower()
        bare = re.sub(r"[^\w-]", "", core)
        if cap_next or bare not in MINOR or len(bare) > 3:
            rendered = "-".join(part[:1].upper() + part[1:] for part in core.split("-"))
        else:
            rendered = core
        out.append(rendered)
        cap_next = bool(re.search(r"[:.?!\u2014]$", word))
    return " ".join(out)


def parse_references(ref_texts: list[str]) -> list[dict]:
    """Pull only fields that are written in the reference line. Never look up a DOI or a date."""
    parsed: list[dict] = []
    for text in ref_texts:
        journal = _parse_narrative_journal(text) or _parse_inverted_journal(text)
        if journal:
            _ground(journal, text)
            parsed.append(journal)
            continue
        page = _parse_webpage(text)
        if page:
            _ground(page, text)
            parsed.append(page)
            continue
        parsed.append({"key": _norm(text)[:24] or "ref", "type": "unparsed", "original": text})
    return parsed


def _ground(ref: dict, original: str) -> None:
    blob = _norm(original)
    for key, value in ref.items():
        if key in {"key", "type"} or value is None:
            continue
        values = [author.get("family", "") for author in value] + [author.get("given") or "" for author in value] if key == "authors" else [value]
        for item in values:
            if item and _norm(str(item)) not in blob:
                raise PagedrillRefusal(
                    "A reference field was not in the author's text. Pagedrill did not invent it.",
                    "grounding_failed",
                )


def _author(family: str, given: str) -> dict:
    return {"family": family, "given": given}


def _parse_narrative_journal(text: str) -> dict | None:
    match = re.match(
        r"^([A-Z][\w'’-]+(?: [A-Z][\w'’-]+)*) and ([A-Z][\w'’-]+(?: [A-Z][\w'’-]+)*) \((\d{4})\)\.\s+\"(.+?)\",\s+(.+?),\s+(\d+)\((\d+)\),\s+pp\.\s+([\d]+-[\d]+)\.\s+doi:\s*(\S+)\s*$",
        text,
    )
    if not match:
        return None
    first, second, year, title, journal, volume, issue, pages, doi = match.groups()
    authors = [_split_given(first), _split_given(second)]
    return {
        "key": authors[0]["family"].lower(),
        "type": "journal-article",
        "authors": authors,
        "year": year,
        "title": title,
        "journal": journal,
        "volume": volume,
        "issue": issue,
        "pages": pages,
        "doi": doi.rstrip("."),
    }


def _split_given(full: str) -> dict:
    parts = full.split(" ")
    return _author(parts[-1], " ".join(parts[:-1]))


def _parse_inverted_journal(text: str) -> dict | None:
    match = re.match(
        r"^((?:[A-Z][\w'’-]+,\s+(?:[A-Z]\.?\s*)+(?:,\s+)?)+)(\d{4})\.\s+(.+)\.\s+(.+?)\s+(\d+),\s+([\d]+-[\d]+)\.\s*$",
        text,
    )
    if not match:
        return None
    author_blob, year, title, journal, volume, pages = match.groups()
    authors = []
    for piece in re.findall(r"([A-Z][\w'’-]+),\s+((?:[A-Z]\.?\s*)+)", author_blob):
        authors.append(_author(piece[0], piece[1].strip()))
    if len(authors) < 1:
        return None
    return {
        "key": authors[0]["family"].lower(),
        "type": "journal-article",
        "authors": authors,
        "year": year,
        "title": title.strip(),
        "journal": journal.strip(),
        "volume": volume,
        "pages": pages,
    }


def _parse_webpage(text: str) -> dict | None:
    match = re.match(r"^([A-Z][\w'’-]+),\s+([A-Z]\.?)\s+(.+)\.\s+([A-Z][\w.&'-]+)\.\s+(https?://\S+)\s*$", text)
    if not match:
        return None
    family, initial, title, site, url = match.groups()
    return {
        "key": family.lower(),
        "type": "webpage-news",
        "authors": [_author(family, initial.rstrip("."))],
        "title": title,
        "site": site,
        "url": url,
    }


def _initials(author: dict) -> str:
    given = (author.get("given") or "").strip()
    if re.fullmatch(r"([A-Z]\.?\s*)+", given):
        letters = re.findall(r"[A-Z]", given)
        return " ".join(f"{letter}." for letter in letters)
    return " ".join(part[0] + "." for part in re.split(r"[\s-]+", given) if part)


def _ref_authors(authors: list[dict]) -> str:
    names = [f'{author["family"]}, {_initials(author)}' for author in authors]
    if len(names) == 1:
        return names[0]
    return ", ".join(names[:-1]) + ", & " + names[-1]


def _en_dash(pages: str) -> str:
    return pages.replace("-", "\u2013")


def apply_user_reference_values(refs: list[dict], supplied: dict | None) -> None:
    """Accept DOI, issue, or date only when the caller passes user_supplied."""
    if not supplied:
        return
    extras = supplied.get("references") or []
    if not isinstance(extras, list):
        return
    by_family = {ref["authors"][0]["family"].lower(): ref for ref in refs if ref.get("type") != "unparsed"}
    for extra in extras:
        if not isinstance(extra, dict):
            continue
        family = str(extra.get("family") or "").lower()
        ref = by_family.get(family)
        if not ref:
            continue
        for key in ("doi", "issue", "date", "year"):
            if extra.get(key):
                ref[key] = str(extra[key]).strip()


def format_apa(original: bytes, profile: dict, supplied: dict | None = None) -> dict:
    supplied = supplied or {}
    source = Document(BytesIO(original))
    paragraphs = [(p.style.name, p.text.strip()) for p in source.paragraphs if p.text.strip()]
    changes: list[dict] = []
    flags: list[dict] = []

    def log(kind: str, before: str, after: str, rule: str) -> None:
        changes.append({"kind": kind, "before": before, "after": after, "rule": rule})

    def flag(item: str, why: str, rule: str) -> None:
        flags.append({"item": item, "why": why, "rule": rule})

    meta: dict[str, str] = {}
    body: list[dict] = []
    ref_texts: list[str] = []
    title = None
    in_refs = False
    for style, text in paragraphs:
        labeled = re.match(r"^(\w+)\s*:\s*(.+)$", text)
        if not in_refs and labeled and labeled.group(1).lower() in LABELS and not body:
            meta[LABELS[labeled.group(1).lower()]] = labeled.group(2).strip()
            continue
        if re.match(r"^(works cited|references?|bibliography|reference list)\s*:?$", text, re.I):
            in_refs = True
            log("label", text, profile["references"]["label"], profile["references"]["src"])
            continue
        if in_refs:
            ref_texts.append(text)
            continue
        if title is None and not style.startswith("Heading") and not body:
            title = text
            continue
        level = int(style.split()[-1]) if style.startswith("Heading") and style.split()[-1].isdigit() else 0
        body.append({"level": level, "text": text})

    if supplied.get("affiliation"):
        meta["affiliation"] = str(supplied["affiliation"]).strip()
    if supplied.get("author_note"):
        meta["author_note"] = str(supplied["author_note"]).strip()
    if supplied.get("course"):
        meta["course"] = str(supplied["course"]).strip()

    professional = profile["title_page"].get("variant") == "professional"
    font_name = profile["font"]["name"]
    font_size = profile["font"]["size_pt"]

    def fix_heading(text: str) -> str:
        stripped = re.sub(r"^\s*(\d+(\.\d+)*[.)]|[a-zA-Z][.)]|[IVXivx]+[.)])\s+", "", text)
        if stripped != text:
            log("heading-numbering", text, stripped, profile["headings"]["src"])
        was_caps = stripped.isupper()
        cased = title_case(stripped)
        if cased != stripped:
            log("heading-case", stripped, cased, profile["headings"]["case_src"])
        if was_caps:
            flag(f'Heading "{text}"', "Was ALL CAPS; converted to title case. Check acronyms and proper nouns.", profile["headings"]["case_src"])
        return cased

    if title:
        cased_title = title_case(title)
        if cased_title != title:
            log("title-case", title, cased_title, profile["headings"]["case_src"])
            title = cased_title
    if body and body[0]["level"] and body[0]["text"].strip().lower().strip(":") in {"introduction", "intro"}:
        log("removed-heading", body[0]["text"], "(removed; paper title acts as Level 1 heading)", profile["headings"]["src"])
        body.pop(0)
    for item in body:
        if item["level"]:
            item["text"] = fix_heading(item["text"])

    for index, item in enumerate(body):
        if not item["level"]:
            continue
        level = item["level"]
        children = []
        for later in body[index + 1 :]:
            if later["level"] and later["level"] <= level:
                break
            if later["level"] == level + 1:
                children.append(later["text"])
        if len(children) == 1:
            flag(
                f'Section "{item["text"]}"',
                f'Has only one Level {level + 1} subsection ("{children[0]}"). APA says to avoid a single subsection. Not changed.',
                profile["headings"]["src"],
            )

    refs = parse_references(ref_texts)
    apply_user_reference_values(refs, supplied)
    originals = {}
    for ref in refs:
        if ref.get("type") == "unparsed":
            flag("Reference", "This reference was not rewritten because its fields could not be read from the line. It is unchanged.", profile["references"]["missing_src"])
            originals[ref["key"]] = ref["original"]
            continue
        hit = next(text for text in ref_texts if _norm(ref["authors"][0]["family"]) in _norm(text))
        originals[ref["key"]] = hit

    rendered = []
    for ref in refs:
        if ref.get("type") == "unparsed":
            rendered.append((ref["key"], [(ref["original"], False, False)], ref))
            continue
        segments = [(_ref_authors(ref["authors"]) + " ", False, False)]
        if ref["type"] == "journal-article":
            segments.append((f'({ref["year"]}). ', False, False))
            title_text = ref["title"].rstrip(".")
            if not ref["title"].endswith("?"):
                title_text += "."
            segments.append((title_text + " ", False, False))
            segments.append((ref["journal"], True, False))
            segments.append((", ", False, False))
            segments.append((ref["volume"], True, False))
            if ref.get("issue"):
                segments.append((f'({ref["issue"]})', False, False))
            else:
                flag(
                    f'Reference {ref["authors"][0]["family"]} ({ref["year"]})',
                    "No issue number in the original. APA includes the issue when the journal uses one. None was added.",
                    "https://apastyle.apa.org/style-grammar-guidelines/references/examples/journal-article-references",
                )
            segments.append((f', {_en_dash(ref["pages"])}.', False, False))
            if ref.get("doi"):
                doi = re.sub(r"^(https?://(dx\.)?doi\.org/|doi:\s*)", "", ref["doi"], flags=re.I)
                segments.append((f" https://doi.org/{doi}", False, False))
                if "doi:" in originals[ref["key"]].lower():
                    log("doi-format", "doi:" + doi, "https://doi.org/" + doi, profile["references"]["doi_src"])
            else:
                flag(
                    f'Reference {ref["authors"][0]["family"]} ({ref["year"]})',
                    "No DOI in the original. If the article has a DOI, APA requires it. None was looked up or added.",
                    profile["references"]["doi_src"],
                )
        elif ref["type"] == "webpage-news":
            if ref.get("date"):
                segments.append((f'({ref["date"]}). ', False, False))
            else:
                segments.append(("(", False, False))
                segments.append(("[DATE NEEDED]", False, True))
                segments.append(("). ", False, False))
                flag(
                    f'Reference {ref["authors"][0]["family"]}',
                    "No publication date in the original. A URL date was not used. Left [DATE NEEDED].",
                    profile["references"]["missing_src"],
                )
            segments.append((ref["title"].rstrip(".") + ".", True, False))
            segments.append((f' {ref["site"]}. {ref["url"]}', False, False))
        rendered.append((ref["authors"][0]["family"].lower(), segments, ref))
        log("reference", originals[ref["key"]], "".join(segment[0] for segment in segments), profile["references"]["elements_src"])

    before_order = [ref["key"] for ref in refs if ref.get("type") != "unparsed"]
    parsed_rendered = [item for item in rendered if item[2].get("type") != "unparsed"]
    unparsed_rendered = [item for item in rendered if item[2].get("type") == "unparsed"]
    parsed_rendered.sort(key=lambda item: item[0])
    if [item[2]["key"] for item in parsed_rendered] != before_order:
        log("reference-order", ", ".join(before_order), ", ".join(item[2]["key"] for item in parsed_rendered), profile["references"]["src"])
    rendered = parsed_rendered + unparsed_rendered

    structured = [ref for ref in refs if ref.get("type") != "unparsed"]
    by_family = {ref["authors"][0]["family"]: ref for ref in structured}

    def year_of(ref: dict) -> str | None:
        return ref.get("year") or ((ref.get("date") or "")[:4] or None)

    def paren_form(ref: dict, year: str) -> str:
        families = [author["family"] for author in ref["authors"]]
        if len(families) == 1:
            names = families[0]
        elif len(families) == 2:
            names = f"{families[0]} & {families[1]}"
        else:
            names = f"{families[0]} et al."
        return f"{names}, {year}"

    def narr_form(ref: dict, year: str) -> str:
        families = [author["family"] for author in ref["authors"]]
        if len(families) == 1:
            names = families[0]
        elif len(families) == 2:
            names = f"{families[0]} and {families[1]}"
        else:
            names = f"{families[0]} et al."
        return f"{names} ({year})"

    def fix_citations(text: str) -> str:
        original_text = text

        def paren(match: re.Match) -> str:
            inner = match.group(1)
            items = [item.strip() for item in inner.split(";")]
            rebuilt = []
            for item in items:
                family = next((name for name in by_family if re.match(rf"^{re.escape(name)}\b", item)), None)
                if not family:
                    return match.group(0)
                year_match = re.search(r"\b(1[89]\d\d|20\d\d|n\.d\.)[a-z]?\b", item)
                if not year_match:
                    flag(f'Citation "({item})"', "No year in the citation. Left as typed.", profile["citations"]["src"])
                    rebuilt.append(item)
                    continue
                rebuilt.append(paren_form(by_family[family], year_match.group(1)))
            if len(rebuilt) > 1:
                rebuilt = sorted(rebuilt, key=str.lower)
            return "(" + "; ".join(rebuilt) + ")"

        text = re.sub(r"\(([^()]*[A-Z][a-z]+[^()]*)\)", paren, text)
        for ref in structured:
            names = [author["family"] for author in ref["authors"]]
            if len(names) >= 3:
                pattern = re.escape(names[0]) + r",\s*" + r",?\s*".join(map(re.escape, names[1:-1])) + r",?\s*(?:and|&)\s*" + re.escape(names[-1])
            elif len(names) == 2:
                pattern = re.escape(names[0]) + r"\s*(?:and|&)\s*" + re.escape(names[1])
            else:
                pattern = re.escape(names[0])
            text = re.sub(pattern + r"\s*\(?\s*((?:19|20)\d\d)\)?(?![,;])", lambda match, ref=ref: narr_form(ref, match.group(1)), text)
        if text != original_text:
            log("in-text-citation", original_text, text, profile["citations"]["src"])
        return text

    before_body = [item["text"] for item in body if not item["level"]]
    for item in body:
        if not item["level"]:
            item["text"] = fix_citations(item["text"])
    after_body = [item["text"] for item in body if not item["level"]]
    surnames = {author["family"] for ref in structured for author in ref["authors"]}
    years = {str(ref.get("year")) for ref in structured if ref.get("year")}
    years |= set(re.findall(r"\b(?:19|20)\d\d\b", " ".join(before_body)))
    assert_apa_body_words("\n".join(before_body), "\n".join(after_body), surnames, years)

    document = Document(BytesIO(original))
    _prepare_styles(document, profile)
    _prepare_section(document, profile)
    _write_header(document, profile, title or "", supplied, flag, log, professional)
    _clear_body(document)
    _write_body(document, profile, title or "", meta, body, rendered, professional, flag, log, bool(supplied.get("table_of_contents")))
    if supplied.get("track_changes"):
        apply_tracked_changes(document, changes)
    fresh = BytesIO()
    document.save(fresh)
    data, verification = splice_allowed(original, fresh.getvalue())
    return {"docx": data, "changes": changes, "flags": flags, "verification": verification, "title": title}


def _prepare_styles(document: Document, profile: dict) -> None:
    font_name = profile["font"]["name"]
    font_size = profile["font"]["size_pt"]
    normal = document.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(font_size)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    paragraph = normal.paragraph_format
    paragraph.line_spacing = profile["line_spacing"]["value"]
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    for level in range(1, 6):
        style = document.styles[f"Heading {level}"]
        spec = profile["headings"][str(level)]
        style.font.name = font_name
        style.font.size = Pt(font_size)
        style.font.bold = spec["bold"]
        style.font.italic = spec["italic"]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            style.element.get_or_add_rPr().append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), font_name)
        heading = style.paragraph_format
        heading.space_before = Pt(0)
        heading.space_after = Pt(0)
        heading.line_spacing = 2.0
        heading.first_line_indent = Inches(0)


def _prepare_section(document: Document, profile: dict) -> None:
    section = document.sections[0]
    section.page_width = Inches(profile["page"]["size_in"][0])
    section.page_height = Inches(profile["page"]["size_in"][1])
    for side, value in profile["page"]["margins_in"].items():
        setattr(section, f"{side}_margin", Inches(value))
    section.header_distance = Inches(0.5)
    section.different_first_page_header_footer = False


def _page_field(paragraph, font_name: str, font_size: int) -> None:
    run = paragraph.add_run()
    run.font.name = font_name
    run.font.size = Pt(font_size)
    for tag, attr in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None), ("w:fldChar", {"w:fldCharType": "separate"})):
        element = OxmlElement(tag)
        if attr:
            for key, value in attr.items():
                element.set(qn(key), value)
        else:
            element.set(qn("xml:space"), "preserve")
            element.text = " PAGE "
        run._r.append(element)
    text = OxmlElement("w:t")
    text.text = "1"
    run._r.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _write_header(document, profile, title: str, supplied: dict, flag, log, professional: bool) -> None:
    header = document.sections[0].header.paragraphs[0]
    previous = header.text
    header.text = ""
    font_name = profile["font"]["name"]
    font_size = profile["font"]["size_pt"]
    if professional and profile["page_header"].get("running_head"):
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), 2)
        running = str(supplied.get("running_head") or "").strip()
        limit = int(profile["page_header"].get("running_head_max_chars") or 50)
        if running:
            if running != running.upper() or len(running) > limit:
                flag("Running head", "The supplied running head must be uppercase and no more than 50 characters. It was not used.", profile["page_header"]["src"])
                running = ""
            else:
                run = header.add_run(running + "\t")
                run.font.name = font_name
                run.font.size = Pt(font_size)
        if not running:
            full = title.upper()
            if len(full) <= limit:
                run = header.add_run(full + "\t")
                run.font.name = font_name
                run.font.size = Pt(font_size)
            else:
                run = header.add_run("[RUNNING HEAD NEEDED]\t")
                run.font.name = font_name
                run.font.size = Pt(font_size)
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                flag(
                    "Running head",
                    "Professional papers use an abbreviated all-caps running head of at most 50 characters. The title is longer, and no shortened title was supplied, so none was invented.",
                    profile["page_header"]["src"],
                )
        _page_field(header, font_name, font_size)
    else:
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _page_field(header, font_name, font_size)
    log("page-header", previous, "[page header]", profile["page_header"]["src"])


def _clear_body(document: Document) -> None:
    body = document.element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def _add_paragraph(document, profile, text="", align="left", bold=False, italic=False, indent=True, style=None, hanging=None, highlight=False):
    paragraph = document.add_paragraph(style=style)
    paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 2.0
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Inches(profile["paragraph"]["first_line_indent_in"]) if indent else Inches(0)
    if hanging:
        fmt.left_indent = Inches(hanging)
        fmt.first_line_indent = Inches(-hanging)
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = profile["font"]["name"]
        run.font.size = Pt(profile["font"]["size_pt"])
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph


def _write_body(document, profile, title, meta, body, rendered, professional, flag, log, table_of_contents: bool = False) -> None:
    title_page = profile["title_page"]
    for _ in range(title_page["blank_lines_before_title"]):
        _add_paragraph(document, profile, indent=False)
    _add_paragraph(document, profile, title, "center", bold=True, indent=False)
    _add_paragraph(document, profile, indent=False)
    if professional:
        _centered_or_placeholder(document, profile, meta.get("author"), "author", title_page["src"], flag)
        _centered_or_placeholder(document, profile, meta.get("affiliation"), "affiliation", title_page["src"], flag)
        _add_paragraph(document, profile, "Author Note", "center", bold=True, indent=False)
        note = meta.get("author_note")
        if note:
            _add_paragraph(document, profile, note, "left", indent=False)
        else:
            _add_paragraph(document, profile, "[AUTHOR NOTE NEEDED]", "left", indent=False, highlight=True)
            flag("Author note", "Not in the original paper. Professional title pages include an author note. None was invented.", title_page["src"])
        for key in ("course", "instructor", "due_date"):
            if meta.get(key):
                log("student-element-omitted", f"{key}: {meta[key]}", "(omitted; professional title page does not include it)", title_page["src"])
        if not any(item["text"].strip().lower() == "abstract" for item in body):
            flag("Abstract", "Professional papers include an abstract on its own page. None was in the original, so none was written.", profile.get("order_src") or title_page["src"])
    else:
        course = meta.get("course")
        if course and not re.match(r"^[A-Z]{2,5}\s?\d{2,4}[A-Z]?:", course):
            updated = re.sub(r"^([A-Z]{2,5}\s?\d{2,4}[A-Z]?)\s+", r"\1: ", course)
            if updated != course:
                log("course-format", course, updated, title_page["src"])
                course = updated
                meta["course"] = course
        if course:
            flag("Course name", f'Kept as typed ("{course}"). Confirm the official course name on the instructional materials.', title_page["src"])
        for key in ("author", "affiliation", "course", "instructor", "due_date"):
            value = meta.get(key)
            if value:
                _add_paragraph(document, profile, value, "center", indent=False)
            else:
                label = "[AFFILIATION NEEDED: Department, Institution]" if key == "affiliation" else f"[{key.upper()} NEEDED]"
                _add_paragraph(document, profile, label, "center", indent=False, highlight=True)
                flag(f"Title page: {key}", "Not in the original paper. Left a highlighted placeholder. Nothing was invented.", title_page["src"])
        for label, key in (("Name:", "author"), ("Professor:", "instructor"), ("Due:", "due_date")):
            if meta.get(key):
                log("title-page-label-removed", f"{label} {meta[key]}", meta[key], title_page["src"])
    document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    if table_of_contents:
        toc = _add_paragraph(document, profile, "", "left", indent=False)
        add_toc_field(toc)
        flag(
            "Table of contents",
            "A Word TOC field was inserted because you asked for one. Word fills it from heading styles. Pagedrill did not type the entries.",
            profile.get("order_src") or profile["title_page"]["src"],
        )
    _add_paragraph(document, profile, title, "center", bold=True, indent=False)
    index = 0
    while index < len(body):
        item = body[index]
        level = item["level"]
        if level:
            spec = profile["headings"][str(level)]
            if spec["inline"]:
                nxt = body[index + 1]["text"] if index + 1 < len(body) and not body[index + 1]["level"] else ""
                paragraph = _add_paragraph(document, profile, "", "left", indent=True, style=f"Heading {level}")
                run = paragraph.add_run(item["text"].rstrip(".") + ".")
                run.bold = True
                run.italic = spec["italic"]
                run.font.name = profile["font"]["name"]
                run.font.size = Pt(profile["font"]["size_pt"])
                if nxt:
                    follow = paragraph.add_run(" " + nxt)
                    follow.bold = False
                    follow.italic = False
                    follow.font.name = profile["font"]["name"]
                    follow.font.size = Pt(profile["font"]["size_pt"])
                index += 2 if nxt else 1
                continue
            align = "center" if spec["align"] == "center" else "left"
            _add_paragraph(document, profile, item["text"], align, bold=spec["bold"], italic=spec["italic"], indent=False, style=f"Heading {level}")
        else:
            _add_paragraph(document, profile, item["text"], "left", indent=True)
        index += 1
    document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    label = profile["references"]["label"] if len(rendered) != 1 else "Reference"
    _add_paragraph(document, profile, label, "center", bold=True, indent=False)
    for _, segments, _ref in rendered:
        paragraph = _add_paragraph(document, profile, "", "left", indent=False, hanging=profile["references"]["hanging_indent_in"])
        for text, italic, highlight in segments:
            run = paragraph.add_run(text)
            run.italic = italic
            run.font.name = profile["font"]["name"]
            run.font.size = Pt(profile["font"]["size_pt"])
            if highlight:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _centered_or_placeholder(document, profile, value, key, rule, flag) -> None:
    if value:
        _add_paragraph(document, profile, value, "center", indent=False)
        return
    label = "[AFFILIATION NEEDED: Department, Institution]" if key == "affiliation" else f"[{key.upper()} NEEDED]"
    _add_paragraph(document, profile, label, "center", indent=False, highlight=True)
    flag(f"Title page: {key}", "Not in the original paper. Left a highlighted placeholder. Nothing was invented.", rule)
