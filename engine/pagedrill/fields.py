"""Word fields. A TOC field is inserted only when the user asks for one."""

from __future__ import annotations

from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_toc_field(paragraph) -> None:
    """Insert a TOC field that Word fills from heading styles. No entries are typed."""
    instruction = ' TOC \\o "1-3" \\h \\z \\u '
    begin = paragraph.add_run()
    char = OxmlElement("w:fldChar")
    char.set(qn("w:fldCharType"), "begin")
    begin._r.append(char)
    instr = paragraph.add_run()
    node = OxmlElement("w:instrText")
    node.set(qn("xml:space"), "preserve")
    node.text = instruction
    instr._r.append(node)
    separate = paragraph.add_run()
    held = OxmlElement("w:fldChar")
    held.set(qn("w:fldCharType"), "separate")
    separate._r.append(held)
    shown = paragraph.add_run(" ")
    shown.font.name = paragraph.style.font.name
    if paragraph.style.font.size is not None:
        shown.font.size = paragraph.style.font.size
    end = paragraph.add_run()
    close = OxmlElement("w:fldChar")
    close.set(qn("w:fldCharType"), "end")
    end._r.append(close)


def is_toc(paragraph) -> bool:
    instruction = "".join(item.text or "" for item in paragraph._p.iter(qn("w:instrText")))
    return "TOC" in instruction
