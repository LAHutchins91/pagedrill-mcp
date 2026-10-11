"""IEEE conference paper, US letter. Journal articles are rendered. Other types stay as typed."""

from __future__ import annotations

import copy
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from .errors import PagedrillRefusal
from .fields import add_toc_field
from .revisions import apply_tracked_changes
from .words import assert_ieee_body_words

JOURNAL = re.compile(
    r"^(?P<authors>.+?) \((?P<year>\d{4}), (?P<month>[A-Za-z]+)\)\. "
    r"(?P<title>.+)\. (?P<container>.+?), "
    r"(?P<volume>\d+)\((?P<issue>[^)]+)\), (?P<locator>\d[\d-]*)\."
    r"(?: (?:https://doi\.org/|doi:)(?P<doi>\S+))?$"
)
CITE_RE = re.compile(r"\(((?:[^();]+?, \d{4}(?:; )?)+)\)")
NARR_RE = re.compile(r"([A-Z][\w\u00C0-\u017F-]+(?: and [A-Z][\w\u00C0-\u017F-]+| et al\.)?) \((\d{4})\)")


def _norm(value: str) -> str:
    return re.sub(r"[\W_]+", "", (value or "").lower())


def _ground(original: str, value: str, label: str) -> None:
    if value and _norm(value) not in _norm(original):
        raise PagedrillRefusal(
            f"A reference {label} was not in the author's entry. Pagedrill did not invent it.",
            "grounding_failed",
        )


def _sentence_case(title: str, keep: set[str]) -> tuple[str, list[str]]:
    out, cap_next, changed = [], True, []

    def acronym(token: str) -> bool:
        return bool(re.search(r"[A-Z].*[A-Z]", token) or re.search(r"\d", token))

    for word in title.split(" "):
        core = re.sub(r"^[^\w]+|[^\w]+$", "", word)
        if core in keep:
            rendered = word
        else:
            parts = []
            for index, part in enumerate(word.split("-")):
                token = re.sub(r"^[^\w]+|[^\w]+$", "", part)
                if (cap_next and index == 0) or acronym(token) or token in keep:
                    parts.append(part)
                else:
                    parts.append(part.lower())
            rendered = "-".join(parts)
        if rendered != word:
            changed.append(word)
        out.append(rendered)
        cap_next = word.endswith(":")
    return " ".join(out), changed


def _headline(title: str, lowercase: set[str]) -> str:
    words = title.split(" ")
    rendered = []
    for index, word in enumerate(words):
        parts = []
        for part in word.split("-"):
            token = part.lower()
            if token in lowercase and 0 < index < len(words) - 1:
                parts.append(token)
            else:
                parts.append(token[:1].upper() + token[1:])
        rendered.append("-".join(parts))
    return " ".join(rendered)


def _authors(text: str) -> list[dict]:
    found = re.findall(r"([^,&]+),\s*((?:[A-Z]\.\s*)+)", text.replace(" & ", ", "))
    return [{"family": family.strip(), "initials": initials.strip()} for family, initials in found]


def _cite_name(authors: list[dict]) -> str:
    if len(authors) == 1:
        return authors[0]["family"]
    if len(authors) == 2:
        return f"{authors[0]['family']} and {authors[1]['family']}"
    return f"{authors[0]['family']} et al."


def parse_journal_line(line: str) -> dict | None:
    match = JOURNAL.match(line.strip())
    if not match:
        return None
    authors = _authors(match.group("authors"))
    if not authors:
        return None
    doi = (match.group("doi") or "").rstrip(".")
    locator = match.group("locator")
    parsed = {
        "type": "journal-article",
        "authors": authors,
        "cite": _cite_name(authors),
        "year": match.group("year"),
        "month": match.group("month"),
        "title": match.group("title").strip(),
        "container": match.group("container").strip(),
        "volume": match.group("volume"),
        "issue": match.group("issue"),
        "pages": locator if "-" in locator else None,
        "article_no": None if "-" in locator else locator,
        "doi": doi or None,
        "typed": line.strip(),
    }
    for author in authors:
        _ground(line, author["family"], "author")
        _ground(line, author["initials"].replace(" ", ""), "initials")
    for key in ("year", "month", "title", "container", "volume", "issue"):
        _ground(line, parsed[key], key)
    if parsed["pages"]:
        _ground(line, parsed["pages"], "pages")
    if parsed["article_no"]:
        _ground(line, parsed["article_no"], "article number")
    if parsed["doi"]:
        _ground(line, parsed["doi"], "doi")
    return parsed


def _pages(value: str) -> str:
    start, end = value.split("-", 1)
    if len(end) < len(start):
        end = start[: len(start) - len(end)] + end
    return f"{start}\u2013{end}"


def _journal_name(name: str, profile: dict, flag, ref: str) -> tuple[str, str]:
    references = profile["references"]
    attested = {}
    attested.update(references.get("guide_attested_titles") or {})
    attested.update(references.get("sheet_attested_titles") or {})
    if name in attested:
        how = "guide example" if name in (references.get("guide_attested_titles") or {}) else "abbreviation sheet title verified for this profile"
        return attested[name], how
    flag(name, f"Journal title left in full. It is not in the attested abbreviation lists ({references.get('sheet_attested_note', 'IEEE sheet')}).", ref)
    return name, "NOT abbreviated"


def render_journal(parsed: dict, profile: dict, log, flag) -> list[tuple[str, bool]]:
    references = profile["references"]
    source = profile["source_keys"]["refguide"]
    authors = parsed["authors"]
    names = [f"{author['initials']} {author['family']}" for author in authors]
    if len(names) > 6:
        who = names[0] + " et al."
    elif len(names) == 1:
        who = names[0]
    elif len(names) == 2:
        who = f"{names[0]} and {names[1]}"
    else:
        who = ", ".join(names[:-1]) + ", and " + names[-1]
    if len(names) == 6:
        flag(parsed["cite"], "Exactly six authors: the Reference Guide lists all six. The conference template wording disagrees, so all six are listed and this is flagged.", source)
    keep = set(references.get("proper_noun_keep") or [])
    title, changed = _sentence_case(parsed["title"], keep)
    if changed:
        log("reference title case", parsed["title"], title, source)
    hits = [word for word in title.split() if re.sub(r"\W", "", word) in keep]
    if changed or hits:
        flag(parsed["cite"] + " title", "Sentence case applied. Proper nouns kept: " + (", ".join(hits) or "none detected") + ". Confirm no other proper noun was lowercased.", source)
    journal, how = _journal_name(parsed["container"], profile, flag, source)
    month = (references.get("months") or {}).get(parsed["month"])
    date = f"{month} {parsed['year']}" if month else parsed["year"]
    if not month:
        flag(parsed["cite"], f"Month '{parsed['month']}' is not in the guide's abbreviation list. The year was kept and the month was not guessed.", source)
    segments = [(f"{who}, \u201c{title},\u201d ", False), (journal, True)]
    tail = f", vol. {parsed['volume']}" + (f", no. {parsed['issue']}" if parsed["issue"] else "")
    if parsed["pages"]:
        pages = _pages(parsed["pages"])
        tail += f", pp. {pages}, {date}"
        if pages.replace("\u2013", "-") != parsed["pages"]:
            log("page range", parsed["pages"], pages, source)
    elif parsed["article_no"]:
        tail += f", {date}, Art. no. {parsed['article_no']}"
    else:
        tail += f", {date}"
        flag(parsed["cite"], "No page range or article number was in the typed entry.", source)
    if parsed["issue"] and "-" in parsed["issue"]:
        flag(parsed["cite"] + " issue", f"Combined issue '{parsed['issue']}' kept as typed. The guide has no rule for issue ranges.", source)
    tail += f", doi: {parsed['doi']}." if parsed["doi"] else "."
    segments.append((tail, False))
    rendered = "".join(text for text, _italic in segments)
    log("reference", parsed["typed"], rendered, source + f" (journal: {how})")
    return segments


def format_ieee(original: bytes, profile: dict, supplied: dict | None = None) -> dict:
    supplied = supplied or {}
    source = Document(BytesIO(original))
    rows = [paragraph for paragraph in source.paragraphs if paragraph.text.strip()]
    if len(rows) < 4:
        raise PagedrillRefusal("This draft does not have a title, author block, and abstract. Nothing was invented.", "missing_field")
    changes: list[dict] = []
    flags: list[dict] = []
    guide = profile["source_keys"]["refguide"]
    esm = profile["source_keys"]["esm"]
    template_src = profile["source_keys"]["template"]

    def log(kind: str, before: str, after: str, rule: str) -> None:
        changes.append({"kind": kind, "before": before, "after": after, "rule": rule})

    def flag(item: str, why: str, rule: str) -> None:
        flags.append({"item": item, "why": why, "rule": rule})

    title = rows[0].text.strip()
    index = 1
    authors = []
    while index < len(rows) and not re.match(r"(?i)abstract\b", rows[index].text):
        authors.append(rows[index].text.strip())
        index += 1
    if index >= len(rows):
        raise PagedrillRefusal("No abstract heading was in the draft. Pagedrill did not invent an abstract.", "missing_field")
    abstract = re.sub(r"(?i)^abstract\s*[:.\u2014-]\s*", "", rows[index].text.strip())
    index += 1
    keywords = None
    if index < len(rows) and re.match(r"(?i)(keywords|index terms)\b", rows[index].text):
        keywords = re.sub(r"(?i)^(keywords|index terms)\s*[:.\u2014-]\s*", "", rows[index].text.strip())
        index += 1
    blocks = []
    while index < len(rows):
        paragraph = rows[index]
        text = paragraph.text.strip()
        if paragraph.style.name.startswith("Heading"):
            label = re.sub(r"^\d+(\.\d+)*\.?\s+", "", text)
            component = (profile["text"].get("component_heads") or {}).get(label.lower())
            if component:
                blocks.append(("comp", component, text))
                if component == "References":
                    index += 1
                    break
            else:
                kind = "h2" if paragraph.style.name == "Heading 2" else "h1"
                blocks.append((kind, label, text))
        else:
            blocks.append(("body", text, text))
        index += 1
    typed_refs = [paragraph.text.strip() for paragraph in rows[index:]]
    parsed_refs = []
    for line in typed_refs:
        parsed = parse_journal_line(line)
        if parsed is None:
            parsed_refs.append({"type": "other", "typed": line, "cite": None, "key": line[:40]})
            flag(line[:80], "This reference is not a journal article. It was kept exactly as typed.", guide)
        else:
            parsed["key"] = parsed["cite"] + parsed["year"]
            parsed_refs.append(parsed)

    order: list[int] = []

    def number_for(name: str, year: str) -> int:
        for ref_index, parsed in enumerate(parsed_refs):
            if parsed.get("cite") == name and parsed.get("year") == year:
                if ref_index not in order:
                    order.append(ref_index)
                return order.index(ref_index) + 1
        raise PagedrillRefusal(
            f"The citation '{name}, {year}' has no matching reference. Pagedrill did not guess one.",
            "grounding_failed",
        )

    new_blocks = []
    for kind, text, original_text in blocks:
        if kind != "body":
            new_blocks.append((kind, text, original_text))
            continue

        def parenthetical(match: re.Match) -> str:
            numbers = sorted({number_for(*piece.rsplit(", ", 1)) for piece in match.group(1).split("; ")})
            return ", ".join(f"[{number}]" for number in numbers)

        def narrative(match: re.Match) -> str:
            number = number_for(match.group(1), match.group(2))
            flag(
                f"{match.group(1)} [{number}]",
                "The author name was kept next to the number. The Editorial Style Manual says to drop it unless it is part of the sentence. Removing words is the author's decision.",
                esm,
            )
            return f"{match.group(1)} [{number}]"

        pieces, cursor = [], 0
        matches = sorted([*CITE_RE.finditer(text), *NARR_RE.finditer(text)], key=lambda item: item.start())
        for match in matches:
            if match.start() < cursor:
                continue
            pieces.append(text[cursor : match.start()])
            pieces.append(parenthetical(match) if match.re is CITE_RE else narrative(match))
            cursor = match.end()
        pieces.append(text[cursor:])
        rendered = "".join(pieces)
        if re.search(r"\breference \[\d+\]", rendered):
            flag("'in reference [n]'", "The Editorial Style Manual says to write 'in [1]', not 'in reference [1]'. The wording was left as the author typed it.", esm)
        if rendered != text:
            log("citation", text, rendered, guide)
        new_blocks.append(("body", rendered, original_text))
    blocks = new_blocks
    missing = [ref_index for ref_index in range(len(parsed_refs)) if ref_index not in order]
    for ref_index in missing:
        flag(parsed_refs[ref_index]["typed"][:80], "Listed but never cited. It is placed last. IEEE numbers by citation order.", template_src)
    order.extend(missing)

    lowercase = set(profile["text"].get("lowercase_words") or [])
    new_title = _headline(title, lowercase) if title.isupper() or title.islower() else title
    if new_title != title:
        log("title case", title, new_title, esm)
        flag("paper title", "The title was typed in all capitals, so acronyms and proper nouns cannot be told apart. Confirm the capitalization.", esm)

    rendered_refs = []
    for ref_index in order:
        parsed = parsed_refs[ref_index]
        if parsed["type"] != "journal-article":
            rendered_refs.append([(parsed["typed"], False)])
            continue
        rendered_refs.append(render_journal(parsed, profile, log, flag))

    body_before = [original_text for kind, _text, original_text in blocks if kind == "body"]
    body_after = [text for kind, text, _original in blocks if kind == "body"]
    surnames = {author["family"] for parsed in parsed_refs if parsed["type"] == "journal-article" for author in parsed["authors"]}
    years = {parsed["year"] for parsed in parsed_refs if parsed.get("year")}
    assert_ieee_body_words("\n".join(body_before), "\n".join(body_after), surnames, years)
    if abstract not in " ".join(paragraph.text for paragraph in rows):
        raise PagedrillRefusal("The abstract text would not match the draft. Pagedrill did not save it.", "word_change")

    template_path = Path(profile["layout"]["base_document"])
    if not template_path.is_file():
        template_path = Path(__file__).resolve().parents[2] / profile["layout"]["base_document"]
    document = Document(str(template_path))
    body = document.element.body
    final_sect = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not final_sect:
            body.remove(child)
    styles = profile["layout"]["styles"]

    def add(style: str, segments: list[tuple[str, bool]]):
        paragraph = document.add_paragraph(style=styles[style])
        for text, italic in segments:
            run = paragraph.add_run(text)
            run.italic = italic
        return paragraph

    add("title", [(new_title, False)])
    for author_index, author in enumerate(authors):
        paragraph = add("author", [(author, False)])
        if author_index:
            paragraph.paragraph_format.space_before = Pt(0)
    if not any("@" in author or "orcid" in author.lower() for author in authors):
        flag("author block", "The template's email or ORCID line is missing. Nothing was invented.", template_src)
    last = document.paragraphs[-1]._p.get_or_add_pPr()
    title_sect = copy.deepcopy(final_sect)
    for cols in title_sect.findall(qn("w:cols")):
        title_sect.remove(cols)
    for kind in title_sect.findall(qn("w:type")):
        title_sect.remove(kind)
    last.append(title_sect)
    cols = final_sect.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        final_sect.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), str(profile["layout"]["body_column_space_twips"]))
    section_type = final_sect.find(qn("w:type"))
    if section_type is None:
        section_type = OxmlElement("w:type")
        final_sect.insert(0, section_type)
    section_type.set(qn("w:val"), "continuous")
    word_count = len(abstract.split())
    add("abstract", [("Abstract", True), ("\u2014" + abstract, False)])
    log("label", "Abstract: ", "Abstract\u2014", template_src)
    if not 150 <= word_count <= 250:
        flag("abstract", f"{word_count} words. The Editorial Style Manual asks for 150-250 words for Transactions. The conference template sets no length. The text was not changed.", esm)
    if keywords:
        add("keywords", [("Keywords\u2014" + keywords, False)])
        log("label", "Keywords: ", "Keywords\u2014", template_src)
    if supplied.get("table_of_contents"):
        toc = document.add_paragraph(style=styles["body"])
        add_toc_field(toc)
        flag("Table of contents", "A Word TOC field was inserted because you asked for one. Entries come from heading styles when Word updates the field. None were typed.", template_src)
    for kind, text, original_text in blocks:
        if kind in ("h1", "h2"):
            add(kind, [(text, False)])
            if text != original_text:
                log("heading number removed", original_text, text, template_src)
        elif kind == "comp":
            add("component_head", [(text, False)])
            if text != original_text:
                log("component heading", original_text, text, esm)
        else:
            add("body", [(text, False)])
    for segments in rendered_refs:
        add("reference", segments)
    flag("footer page numbers", "Pagination is omitted. The template says not to add pagination anywhere in the paper.", template_src)
    if supplied.get("track_changes"):
        apply_tracked_changes(document, changes)
    buffer = BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    return {
        "docx": data,
        "changes": changes,
        "flags": flags,
        "verification": {
            "changed_parts": ["word/document.xml", "word/styles.xml", "word/numbering.xml"],
            "removed_parts": [],
            "added_parts": [],
            "untouched_part_count": 0,
            "untouched_identical": True,
            "basis": "ieee-conference-template",
        },
        "title": new_title,
        "notice": None,
    }
