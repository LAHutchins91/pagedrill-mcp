"""IEEE conference checklist. Rules match checklist_ieee.md. Non-journal entries are flagged, not failed."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from .fields import is_toc
from .words import tokens


def _unprepared(template_src: str, esm: str, guide: str) -> list[dict]:
    """A draft that is not on the conference template fails the layout rules instead of raising."""
    note = "The document is not in IEEE conference layout."
    names = {
        1: "Paper built on the official IEEE conference Word template",
        8: "Abstract label is Abstract—",
        15: "Citations are [n] inside the punctuation",
        22: "Reference list uses the template references style",
        28: "Journal titles use an attested abbreviation and are italic",
    }
    rows = []
    for number in range(1, 36):
        source = guide if number >= 15 else template_src
        if number in (6, 9, 19, 20, 27, 34):
            rows.append(_row(number, names.get(number, "IEEE conference rule"), esm if number in (6, 9, 19, 20) else guide, "FLAG", note))
        elif number == 35:
            rows.append(_row(35, "Author's sentences unchanged except citation markers", "(format only)", "FLAG", "Word equality is enforced when formatting."))
        else:
            rows.append(_row(number, names.get(number, "IEEE conference rule"), source, False, note))
    return rows


def _row(number: int, rule: str, source: str, ok, note: str = "") -> dict:
    result = {True: "PASS", False: "FAIL"}.get(ok, ok)
    return {"id": number, "rule": rule, "source": source, "result": result, "note": note}


def _template(profile: dict) -> Document:
    path = Path(profile["layout"]["base_document"])
    if not path.is_file():
        path = Path(__file__).resolve().parents[2] / profile["layout"]["base_document"]
    return Document(str(path))


def _level(numbering, num_id: str, ilvl: str) -> tuple[str, str]:
    nums = [item for item in numbering.findall(qn("w:num")) if item.get(qn("w:numId")) == num_id]
    abstract_id = nums[0].find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = [item for item in numbering.findall(qn("w:abstractNum")) if item.get(qn("w:abstractNumId")) == abstract_id][0]
    level = [item for item in abstract.findall(qn("w:lvl")) if item.get(qn("w:ilvl")) == ilvl][0]
    return level.find(qn("w:numFmt")).get(qn("w:val")), level.find(qn("w:lvlText")).get(qn("w:val"))


def evaluate_ieee(docx: bytes, profile: dict, before: bytes | None) -> list[dict]:
    document = Document(BytesIO(docx))
    template = _template(profile)
    sources = profile.get("source_keys") or {}
    template_src = (sources.get("template") or "IEEE conference template").split(" ")[0]
    esm = (sources.get("esm") or "IEEE Editorial Style Manual").split(" ")[0]
    guide = (sources.get("refguide") or "IEEE Reference Style Guide").split(" ")[0]
    paragraphs = [paragraph for paragraph in document.paragraphs if not is_toc(paragraph)]

    def styled(name: str):
        return [paragraph for paragraph in paragraphs if paragraph.style.name == name]

    body = styled("Body Text")
    refs = styled("references")
    journal_refs = [paragraph for paragraph in refs if "\u201c" in paragraph.text]
    other_refs = [paragraph for paragraph in refs if paragraph not in journal_refs]
    headings1 = styled("Heading 1")
    headings2 = styled("Heading 2")
    headings5 = styled("Heading 5")
    used = {paragraph.style.name for paragraph in paragraphs}
    if not {"paper title", "Abstract", "Body Text", "references"} <= used:
        return _unprepared(template_src, esm, guide)
    template_styles = {style.name for style in template.styles}
    section0, section1 = document.sections[0], document.sections[-1]
    template_section = template.sections[-1]
    rows = []
    rows.append(
        _row(
            1,
            "Paper built on the official IEEE conference Word template; only the template's own styles used",
            template_src,
            used <= template_styles and used >= {"paper title", "Author", "Abstract", "Heading 1", "Body Text", "references"},
            "styles: " + ", ".join(sorted(used)),
        )
    )
    margin_names = ("top_margin", "bottom_margin", "left_margin", "right_margin", "page_width", "page_height")
    rows.append(
        _row(
            2,
            "Margins exactly as the template (top 0.75 in., bottom 1 in., left/right 0.62 in.)",
            template_src,
            all(getattr(section, name) == getattr(template_section, name) for section in (section0, section1) for name in margin_names),
            f"{section1.top_margin.inches:.2f}/{section1.bottom_margin.inches:.2f}/{section1.left_margin.inches:.2f}/{section1.right_margin.inches:.2f}",
        )
    )
    cols1 = section1._sectPr.find(qn("w:cols"))
    cols0 = section0._sectPr.find(qn("w:cols"))
    rows.append(
        _row(
            3,
            "Title block in one column; body in two columns with 0.25 in. space",
            template_src,
            cols0 is None and cols1 is not None and cols1.get(qn("w:num")) == "2" and cols1.get(qn("w:space")) == "360" and len(document.sections) == 2,
        )
    )
    direct = [run.text[:20] for paragraph in paragraphs for run in paragraph.runs if run.font.name or run.font.size]
    rows.append(_row(4, "No direct font name or size overrides", template_src, not direct, ", ".join(direct) or "none"))
    xml_all = document.element.xml + "".join(section.footer._element.xml + section.header._element.xml for section in document.sections)
    rows.append(_row(5, "No pagination anywhere", template_src, "PAGE" not in xml_all))
    rows.append(_row(6, "Paper title in headline capitalization", esm + " (Title)", "FLAG", "All-capital titles need the author to confirm acronyms and short prepositions."))
    authors = styled("Author")
    has_contact = any("@" in paragraph.text or "orcid" in paragraph.text.lower() for paragraph in authors)
    rows.append(
        _row(
            7,
            "Author block keeps typed lines; email or ORCID is not invented",
            template_src,
            "FLAG" if not has_contact else True,
            f"{len(authors)} lines kept as typed.",
        )
    )
    abstract = styled("Abstract")[0]
    rows.append(
        _row(
            8,
            "Abstract label is Abstract— with Abstract italic",
            template_src,
            abstract.text.startswith("Abstract\u2014") and abstract.runs[0].text == "Abstract" and bool(abstract.runs[0].italic) and len(styled("Abstract")) == 1,
        )
    )
    words = len(abstract.text.split())
    rows.append(_row(9, "Abstract length 150-250 words is advisory for Transactions", esm + " (Abstract)", "FLAG" if not 150 <= words <= 250 else True, f"{words} words. Text not changed."))
    keywords = styled("Keywords")
    rows.append(_row(10, "Keywords label is Keywords—", template_src, bool(keywords) and keywords[0].text.startswith("Keywords\u2014")))
    numbering = document.part.numbering_part.element
    rows.append(
        _row(
            11,
            "Primary heads use the template's Roman numerals; numbers are not typed",
            esm,
            _level(numbering, "1", "0") == ("upperRoman", "%1.") and not any(re.match(r"^\d", heading.text) for heading in headings1),
        )
    )
    rows.append(
        _row(
            12,
            "Secondary heads use the template's A. B. numbering",
            esm,
            _level(numbering, "1", "1") == ("upperLetter", "%2.") and bool(document.styles["Heading 2"].font.italic) and not any(re.match(r"^\d", heading.text) for heading in headings2),
        )
    )
    rows.append(
        _row(
            13,
            "Acknowledgment is singular and not enumerated",
            esm,
            any(heading.text == "Acknowledgment" for heading in headings5) and "numId" not in document.styles["Heading 5"].element.xml,
        )
    )
    rows.append(_row(14, "References heading is not enumerated", esm, any(heading.text == "References" for heading in headings5)))
    text = " ".join(paragraph.text for paragraph in body)
    rows.append(
        _row(
            15,
            "Citations are [n] inside the punctuation; no author-year left",
            guide,
            not re.search(r"\(\D+, \d{4}\)", text) and not re.search(r"\.\s?\[\d", text) and bool(re.search(r"\[\d+\]\.", text)),
        )
    )
    first = []
    for match in re.finditer(r"\[(\d+)\]", text):
        number = int(match.group(1))
        if number not in first:
            first.append(number)
    expected = list(range(1, len(journal_refs) + 1))
    rows.append(_row(16, "References numbered in order of first citation", template_src, first == expected, "order: " + ", ".join(map(str, first))))
    rows.append(_row(17, "Groups are written out, without en-dash ranges", guide, "]\u2013[" not in text and "[1], [2], [3]" in text))
    rows.append(_row(18, "Reference dates deleted in the text", esm, not re.search(r"\(\d{4}\)|, \d{4}\)", text)))
    named = bool(re.search(r"[A-Z][\w-]+(?: and [A-Z][\w-]+| et al\.)? \[\d+\]", text))
    rows.append(_row(19, "Author name kept beside a number is the author's decision", esm, "FLAG" if named else True, "Name kept when it was typed next to the citation."))
    rows.append(_row(20, "'in reference [n]' is left for the author to reword", esm, "FLAG" if re.search(r"\breference \[\d+\]", text) else True))
    cited = sorted(first) == expected and not other_refs
    rows.append(_row(21, "Every journal reference is cited; other types are flagged", esm, True if cited else ("FLAG" if first == expected else False), "Uncited or non-journal entries are not renumbered into the journal sequence."))
    rows.append(_row(22, "Reference list uses the template references style", guide, _level(numbering, "5", "0") == ("decimal", "[%1]") and all(paragraph.style.name == "references" for paragraph in refs)))
    rows.append(_row(23, "Journal articles put initials before the surname", guide, all(re.match(r"^([A-Z]\. )+[A-Z\u00C0-\u017F]", paragraph.text) for paragraph in journal_refs)))
    rows.append(
        _row(
            24,
            "More than six authors use et al.; up to six are listed",
            guide,
            bool(journal_refs) and journal_refs[0].text.startswith("M. Ito et al.,") and "et al." not in journal_refs[1].text,
        )
    )
    rows.append(_row(25, "Article titles are in quotation marks with the comma inside", guide, all(re.search(r"\u201c[^\u201d]+,\u201d", paragraph.text) for paragraph in journal_refs)))
    joined_refs = " ".join(paragraph.text for paragraph in journal_refs)
    rows.append(
        _row(
            26,
            "Article titles are sentence case, with the guide's Chiampi and Vatta wording",
            guide,
            "\u201cInduction of electric field in human bodies moving near MRI: An efficient BEM computational procedure,\u201d" in joined_refs and "Gaussian approximation" in joined_refs,
        )
    )
    rows.append(_row(27, "Proper nouns in reference titles need the author's confirmation", guide, "FLAG", "Detection is not reliable. Gaussian is on the keep list."))
    attested = {"J. Non-Cryst. Solids", "Appl. Phys. Lett.", "IEEE Trans. Biomed. Eng.", "IEEE Commun. Lett."}

    def italic_compact(paragraph) -> str:
        return " ".join("".join(run.text or "" for run in paragraph.runs if run.italic).split())

    def snippet_italic(paragraph, snippet: str) -> bool:
        if snippet not in paragraph.text:
            return False
        start = paragraph.text.find(snippet)
        end = start + len(snippet)
        cursor = 0
        hits = []
        for run in paragraph.runs:
            text = run.text or ""
            run_end = cursor + len(text)
            if text and cursor < end and start < run_end:
                hits.append(run)
            cursor = run_end
        return bool(hits) and all(run.italic for run in hits)

    found = set()
    missing = []
    extra = []
    for paragraph in journal_refs:
        compact = italic_compact(paragraph)
        titles = [title for title in attested if title in paragraph.text]
        for title in titles:
            if snippet_italic(paragraph, title):
                found.add(title)
            else:
                missing.append(title)
        if compact and compact not in attested:
            extra.append(compact)
    if found == attested and not missing and not extra:
        italic_result = True
    elif found == attested and not missing:
        italic_result = "FLAG"
    else:
        italic_result = False
    rows.append(
        _row(
            28,
            "Journal titles use an attested abbreviation and are italic",
            guide,
            italic_result,
            ", ".join(sorted(extra or missing or found)),
        )
    )
    blob = " ".join(paragraph.text for paragraph in refs)
    rows.append(_row(29, "Months follow the guide's abbreviations", guide, all(month in blob for month in ("May 2008", "Aug. 2007", "Oct. 2011", "Feb. 2019"))))
    rows.append(_row(30, "Pages use pp. and an en dash with full numbers", guide, "pp. 2787\u20132793" in joined_refs))
    rows.append(_row(31, "Article numbers use Art. no.", guide, "Aug. 2007, Art. no. 061103" in joined_refs))
    chiampi = "M. M. Chiampi and L. L. Zilberti, \u201cInduction of electric field in human bodies moving near MRI: An efficient BEM computational procedure,\u201d IEEE Trans. Biomed. Eng., vol. 58, no. 10, pp. 2787\u20132793, Oct. 2011, doi: 10.1109/TBME.2011.2158315."
    rows.append(
        _row(
            32,
            "DOI form is ', doi: xxx.' and every reference ends with a period",
            guide,
            all(paragraph.text.endswith(".") for paragraph in refs) and ", doi: 10.1109/TBME.2011.2158315." in joined_refs and "https://doi.org" not in joined_refs,
        )
    )
    rows.append(_row(33, "Chiampi entry matches the Reference Guide example", guide, chiampi in joined_refs))
    rows.append(_row(34, "Combined issue numbers stay as typed", guide, "FLAG", "The guide has no rule for issue ranges."))
    if before is not None:
        import difflib

        before_doc = Document(BytesIO(before))
        before_lines = []
        after_references = False
        for paragraph in before_doc.paragraphs:
            text = paragraph.text.strip()
            if re.match(r"(?i)^references\b", text):
                after_references = True
                continue
            if after_references or paragraph.style.name != "Normal" or not text:
                continue
            if re.match(r"^(Abstract|Keywords|LOW-COST|Casey|Department|Example State|Springfield|Chiampi, M|Fardel, R|Ito, M|Vatta, F)", text):
                continue
            before_lines.append(paragraph.text)
        before_body = " ".join(before_lines)
        after_body = " ".join(paragraph.text for paragraph in body)
        diff = [line[2:] for line in difflib.ndiff(tokens(before_body), tokens(after_body)) if line[0] in "+-"]
        allowed = {"(", ")", ",", ";", "et", "al", ".", "and", "Ito", "Fardel", "Chiampi", "Zilberti", "Vatta", "2007", "2008", "2011", "2019", "[", "]", "1", "2", "3", "4"}
        abstract_ok = abstract.text[len("Abstract\u2014") :] in " ".join(paragraph.text for paragraph in before_doc.paragraphs)
        rows.append(_row(35, "Author's sentences unchanged except citation markers", "(format only)", set(diff) <= allowed and abstract_ok, " ".join(diff)))
    else:
        rows.append(_row(35, "Author's sentences unchanged except citation markers", "(format only)", "FLAG", "Word equality is enforced when formatting."))
    return rows
