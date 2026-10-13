"""Vancouver checklist. A rule passes only when the document shows it.

Journal options (grouped citation punctuation, list numbering, page layout) are
FLAG: ICMJE and Citing Medicine leave them to the journal.
"""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document

from .errors import PagedrillRefusal
from .vancouver import CH1, CH2, OFFICIAL, PAGE_EXAMPLES, elide_pages
from .words import assert_vancouver_body_words

ICMJE = "https://www.icmje.org/recommendations/browse/manuscript-preparation/preparing-for-submission.html"
NLM = "https://www.nlm.nih.gov/bsd/uniform_requirements.html"
MONTHS = "Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec"


def evaluate_vancouver(docx: bytes, profile: dict, before: bytes | None) -> list[dict]:
    document = Document(BytesIO(docx))
    paragraphs = list(document.paragraphs)
    ref_at = next((index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip().lower() in {"references", "reference list", "bibliography", "works cited"}), None)
    if ref_at is None:
        return [_row(number, "Vancouver reference list", ICMJE, False, "No References heading.") for number in range(1, 18)]
    body = " ".join(paragraph.text for paragraph in paragraphs[:ref_at])
    refs = [paragraph.text.strip() for paragraph in paragraphs[ref_at + 1 :] if paragraph.text.strip()]
    stripped = [_strip(line) for line in refs]
    got = {}
    for line in stripped:
        name = line.split(" ")[0].strip(",.;") if line else ""
        if name:
            got[name] = line
    rows = [
        _citations(body),
        _order(body, refs),
        _multiple(body),
        _row(4, "Placement of the citation number relative to punctuation", ICMJE, "FLAG", "Left where the author typed it. ICMJE does not specify."),
        _cited(body, refs),
        _numbering(refs),
        _samples(got),
        _authors(stripped),
        _abbreviations(got, profile),
        _titles(stripped),
        _dates(stripped),
        _pages(stripped),
        _doi(got),
        _books(got),
        _layout(docx, before),
        _row(
            16,
            "Reference data checked against PubMed/Crossref",
            "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi",
            "FLAG",
            "Halpern, Zhang, and Rose were verified in PubMed when the sample was built. Rose also exists as PMID 12088834. Murray is a book and was not looked up. Nothing was fetched at format time.",
        ),
        _words(docx, before),
    ]
    return rows


def _row(number: int, rule: str, source: str, ok, note: str = "") -> dict:
    result = {True: "PASS", False: "FAIL"}.get(ok, ok)
    return {"id": number, "rule": rule, "source": source, "result": result, "note": note}


def _strip(line: str) -> str:
    return re.sub(r"^(?:\(\d+\)\s+|\d+\.\s+|\d+\s+)", "", line)


def _number_groups(body: str) -> list[str]:
    return re.findall(r"\(([\d][\d,\- ]*)\)", body)


def _citations(body: str) -> dict:
    author_year = re.search(r"\([A-Za-z][^)]{0,80}\d{4}\)", body)
    groups = _number_groups(body)
    shape = re.compile(r"\d+(?:-\d+)?(?:, ?\d+(?:-\d+)?)*")
    ok = author_year is None and bool(groups) and all(shape.fullmatch(group.strip()) for group in groups)
    return _row(1, "Citations identified by Arabic numerals in parentheses; no author-year left", ICMJE + " (g. References)", ok)


def _order(body: str, refs: list[str]) -> dict:
    first: list[int] = []
    for group in _number_groups(body):
        for piece in re.split(r",\s*", group):
            for number in re.findall(r"\d+", piece):
                value = int(number)
                if value not in first:
                    first.append(value)
    ok = bool(first) and first == list(range(1, len(first) + 1)) and max(first) <= len(refs)
    return _row(2, "References numbered consecutively in order of first mention", ICMJE + " (g. References)", ok, "first-mention order: " + ", ".join(map(str, first)))


def _multiple(body: str) -> dict:
    groups = [group for group in _number_groups(body) if "," in group or "-" in group]
    if not groups:
        return _row(3, "Several references in one citation: separator/range style", ICMJE, "FLAG", "No grouped citation in this document. The separator is a journal option, not an ICMJE rule.")
    return _row(
        3,
        "Several references in one citation: separator/range style",
        ICMJE,
        "FLAG",
        "Used " + ", ".join(f"({group.strip()})" for group in groups) + ". ICMJE does not specify commas, spaces, or ranges. Citing Medicine chapters 1 and 2 do not set in-text punctuation. The profile records the journal option.",
    )


def _cited(body: str, refs: list[str]) -> dict:
    seen = {int(number) for group in _number_groups(body) for number in re.findall(r"\d+", group)}
    expected = set(range(1, len(refs) + 1))
    if seen == expected:
        return _row(5, "Every listed reference cited; one reference per number", ICMJE, True)
    if expected - seen and max(expected - seen) == len(refs):
        return _row(5, "Every listed reference cited; one reference per number", ICMJE, "FLAG", "A listed reference is not cited in the text. It was not dropped.")
    return _row(5, "Every listed reference cited; one reference per number", ICMJE, False, "Citation numbers and the reference list do not match.")


def _numbering(refs: list[str]) -> dict:
    if refs and all(re.match(r"^\d+\.\s+\S", line) for line in refs):
        used = "'1. '"
    elif refs and all(re.match(r"^\(\d+\)\s+\S", line) for line in refs):
        used = "'(1) '"
    elif refs and all(re.match(r"^\d+\s+\S", line) for line in refs):
        used = "'1 '"
    else:
        return _row(6, "Reference list numbering format", ICMJE + "; " + NLM, False, "The list is not numbered with one of the journal options ('1. ', '1 ', or '(1) ').")
    return _row(
        6,
        "Reference list numbering format",
        ICMJE + "; " + NLM,
        "FLAG",
        f"{used} is the selected journal option. Neither ICMJE nor the Sample References page shows the list number. Citing Medicine examples are unnumbered component strings.",
    )


def _samples(got: dict[str, str]) -> dict:
    present = [name for name in OFFICIAL if name in got]
    ok = all(got[name] == OFFICIAL[name] for name in present)
    note = "; ".join(f"{name}: {'match' if got.get(name) == OFFICIAL[name] else 'DIFF'}" for name in present) or "No NLM sample reference is in this list."
    return _row(7, "Each known sample entry identical to NLM's Sample Reference for that work", NLM, ok, note)


def _authors(stripped: list[str]) -> dict:
    ok = True
    for line in stripped:
        if not _journal(line):
            continue
        head = line.split(". ")[0]
        if head.endswith("et al"):
            names = [name for name in head[: -len("et al")].strip(" ,").split(", ") if name]
            if len(names) != 6:
                ok = False
        else:
            names = [name for name in head.split(", ") if name]
            if len(names) > 6:
                ok = False
    rose = next((line for line in stripped if line.startswith("Rose ")), "")
    zhang = next((line for line in stripped if line.startswith("Zhang ")), "")
    if rose and not rose.startswith("Rose ME, Huerbin MB, Melick J, Marion DW, Palmer AM, Schiding JK, et al."):
        ok = False
    if zhang and "et al" in zhang.split(". ")[0]:
        ok = False
    return _row(
        8,
        "Journal authors: surname plus initials; first six, then et al. when more than six. Books list every author.",
        NLM + " (#1); " + CH1 + ", Author, General Rules and example 3 (optional limit to 6); " + CH2 + ", Author/Editor, General Rules (give all)",
        ok,
        "Appendix F (MEDLINE notes) is not in the chapter PDFs. The six-author journal limit is the NLM Sample References sentence and Citing Medicine's own 'to 6 authors' example.",
    )


def _abbreviations(got: dict[str, str], profile: dict) -> dict:
    expect = {"Halpern": "N Engl J Med.", "Rose": "Brain Res.", "Zhang": "BMJ."}
    ok = all(token in got[name] for name, token in expect.items() if name in got)
    full = ("New England Journal of Medicine", "Brain Research")
    if any(name in got and title in got[name] for name, title in (("Halpern", full[0]), ("Rose", full[1]))):
        ok = False
    note = (profile.get("references") or {}).get("medline_attested_note", "Attested MEDLINE abbreviations only.")
    return _row(9, "Journal titles abbreviated per an attested MEDLINE form, not guessed", ICMJE + " (MEDLINE abbreviation)", ok, note)


def _titles(stripped: list[str]) -> dict:
    titles = []
    for line in stripped:
        parts = line.split(". ")
        if len(parts) >= 2:
            titles.append(parts[1])
    ok = bool(titles) and all(_title_ok(title) for title in titles)
    return _row(
        10,
        "Titles in sentence case; lowercase after a colon; acronyms kept",
        CH1 + ", Article Title for Journal Articles, General Rules; " + CH2 + ", Title for Entire Books, General Rules",
        ok,
        "Capitalize only the first word, proper nouns, proper adjectives, acronyms, and initialisms. Proper nouns are flagged for the author to confirm.",
    )


def _title_ok(title: str) -> bool:
    words = title.split(" ")
    for index, word in enumerate(words):
        for part_index, part in enumerate(word.split("-")):
            core = re.sub(r"[^\w]", "", part, flags=re.UNICODE)
            if not core:
                continue
            if sum(char.isupper() for char in core) >= 2 or any(char.isdigit() for char in core):
                continue
            if index == 0 and part_index == 0:
                continue
            if any(char.isupper() for char in core):
                return False
    return True


def _dates(stripped: list[str]) -> dict:
    journals = [line for line in stripped if _journal(line)]
    pattern = re.compile(rf"\d{{4}}(?: (?:{MONTHS})(?: \d{{1,2}})?)?[;:]")
    period = re.compile(rf"\b(?:{MONTHS})\.")
    ok = bool(journals) and all(pattern.search(line) and not period.search(line) for line in journals)
    return _row(
        11,
        "Date as year, then a three-letter month and day when the entry has them",
        CH1 + ", Date of Publication for Journal Articles, General Rules",
        ok,
        "Example in the chapter: 2004 May 5. A missing month or day is omitted, not invented.",
    )


def _pages(stripped: list[str]) -> dict:
    examples_ok = all(elide_pages(source) == target for source, target in PAGE_EXAMPLES.items())
    shown = True
    for line in stripped:
        if not _journal(line):
            continue
        for pages in re.findall(r":([A-Za-z0-9,\-]+)\.", line):
            if elide_pages(pages) != pages:
                shown = False
    sample = ("284-7" in " ".join(stripped) or "Halpern" not in " ".join(stripped)) and ("40-6" in " ".join(stripped) or "Rose" not in " ".join(stripped))
    return _row(
        12,
        "Page ranges shortened; a following letter or a roman numeral is not shortened",
        CH1 + ", Location (Pagination) for Journal Articles, General Rules; examples 59-61",
        examples_ok and shown and sample,
        "123-125 becomes 123-5. 124A-126A stays. A leading letter stays (S10-8). Roman numerals stay as found (xvii-xix). Discontinuous pages shorten each span (24-29 becomes 24-9).",
    )


def _doi(got: dict[str, str]) -> dict:
    ok = "Zhang" not in got or got["Zhang"].endswith("doi: 10.1136/bmj.a2752.")
    if any("doi:" in line and not re.search(r" doi: 10\.\S+\.$", line) for line in got.values()):
        ok = False
    return _row(13, "DOI after the pagination as 'doi: 10.xxxx.'", NLM + " (Article with a Digital Object Identifier)", ok)


def _books(got: dict[str, str]) -> dict:
    ok = "Murray" not in got or got["Murray"] == OFFICIAL["Murray"]
    return _row(
        14,
        "Entire book: Authors. Title. Edition. Place: Publisher; year.",
        CH2 + ", Chapter 2A Entire Books (Author/Editor, Title, Edition, Place of Publication, Publisher, Date of Publication); NLM sample #19",
        ok,
        "Authors are all listed. Edition uses an arabic ordinal. Place and publisher stay as typed. Chapter 2C page shortening is the same rule as journal pages; this sample has no part of a book.",
    )


def _layout(docx: bytes, before: bytes | None) -> dict:
    if before is None:
        return _row(15, "Page layout (font, spacing, margins, title page)", ICMJE, "FLAG", "ICMJE sets none. The journal option leaves the author's layout unchanged.")
    current = _layout_signature(docx)
    previous = _layout_signature(before)
    if current != previous:
        return _row(15, "Page layout (font, spacing, margins, title page)", ICMJE, False, "Font or margins changed. The journal option is to leave them unchanged.")
    return _row(15, "Page layout (font, spacing, margins, title page)", ICMJE, "FLAG", "Left unchanged. ICMJE says to consult the journal. No font or margin was guessed.")


def _layout_signature(data: bytes) -> tuple:
    document = Document(BytesIO(data))
    section = document.sections[0]
    font = document.styles["Normal"].font
    size = None if font.size is None else font.size.pt
    return (
        section.top_margin,
        section.bottom_margin,
        section.left_margin,
        section.right_margin,
        section.page_width,
        section.page_height,
        font.name,
        size,
    )


def _words(docx: bytes, before: bytes | None) -> dict:
    if before is None:
        return _row(17, "Author's sentences unchanged (only citation markers change)", "(format only)", "FLAG", "Word equality is enforced when formatting.")
    current = _body(docx)
    previous = _body(before)
    surnames = set(re.findall(r"\b[A-Z][a-zA-Z'’\-]+\b", previous))
    years = set(re.findall(r"\b(?:19|20)\d\d\b", previous))
    try:
        assert_vancouver_body_words(previous, current, surnames, years)
    except PagedrillRefusal:
        return _row(17, "Author's sentences unchanged (only citation markers change)", "(format only)", False, "The body wording changed beyond citation markers.")
    return _row(17, "Author's sentences unchanged (only citation markers change)", "(format only)", True)


def _body(data: bytes) -> str:
    document = Document(BytesIO(data))
    paragraphs = list(document.paragraphs)
    ref_at = next((index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip().lower() in {"references", "reference list", "bibliography", "works cited"}), len(paragraphs))
    return "\n".join(paragraph.text for paragraph in paragraphs[:ref_at])


def _journal(line: str) -> bool:
    return bool(re.search(rf"\d{{4}}(?: (?:{MONTHS})(?: \d{{1,2}})?)?;", line))
